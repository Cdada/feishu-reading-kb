#!/usr/bin/env python3
from __future__ import annotations

import argparse
import base64
import csv
import datetime as dt
import hashlib
import json
import os
import re
import signal
import shutil
import subprocess
import sys
import urllib.parse
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from webarchive_converter import convert_webarchive_to_markdown


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = PROJECT_ROOT / "config" / "defaults.json"

MANIFEST_FIELDS = [
    "batch_id",
    "original_path",
    "relative_input_path",
    "group_id",
    "group_role",
    "group_main",
    "linked_from",
    "link_target",
    "staged_path",
    "file_name",
    "extension",
    "size_bytes",
    "modified_time",
    "file_kind",
    "action",
    "processing_status",
    "kb_folder",
    "classification",
    "primary_input",
    "auxiliary_input",
    "quality_status",
    "source_url",
    "conversion_method",
    "conversion_status",
    "notes",
]

KB_MANIFEST_FIELDS = [
    "source_file",
    "group_id",
    "group_role",
    "linked_from",
    "classification",
    "selected_for_kb",
    "primary_input",
    "auxiliary_input",
    "needs_manual_review",
    "source_url",
    "conversion_method",
    "conversion_status",
    "notes",
]

MINERU_EXTENSIONS = {
    ".pdf",
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp",
}
HEIC_EXTENSIONS = {".heic", ".heif"}
DOCX_EXTENSIONS = {".docx"}
LEGACY_DOC_EXTENSIONS = {".doc"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}
TEXT_EXTENSIONS = {".md", ".txt", ".csv"}
PRESENTATION_EXTENSIONS = {".pptx", ".ppt"}
WEBARCHIVE_EXTENSIONS = {".webarchive"}
IMAGE_ZIP_EXTENSIONS = {".zip"}
SUPPORTED_EXTENSIONS = (
    MINERU_EXTENSIONS
    | HEIC_EXTENSIONS
    | DOCX_EXTENSIONS
    | LEGACY_DOC_EXTENSIONS
    | SPREADSHEET_EXTENSIONS
    | TEXT_EXTENSIONS
    | PRESENTATION_EXTENSIONS
    | WEBARCHIVE_EXTENSIONS
    | IMAGE_ZIP_EXTENSIONS
)

PROCESS_FILE_NAMES = {
    "middle.json",
    "model.json",
    "layout.pdf",
    "span.pdf",
}
JUNK_FILE_NAMES = {
    ".DS_Store",
}
MINERU_BACKENDS = {"pipeline"}
MINERU_PARSE_METHODS = {"auto", "txt", "ocr"}
MINERU_LANGUAGES = {
    "ch",
    "ch_server",
    "ch_lite",
    "en",
    "korean",
    "japan",
    "chinese_cht",
    "ta",
    "te",
    "ka",
    "th",
    "el",
    "latin",
    "arabic",
    "east_slavic",
    "cyrillic",
    "devanagari",
}


@dataclass(frozen=True)
class Config:
    icloud_inbox: Path
    local_transfer_station: Path
    conversion_tool_workspace: Path
    result_prefix: str
    default_topic: str
    markdown_preview_rows_per_sheet: int
    skip_name_prefixes: tuple[str, ...]
    skip_name_contains: tuple[str, ...]
    mineru_backend: str = "pipeline"
    mineru_parse_method: str = "auto"
    mineru_language: str = "ch"
    preserve_webarchive_images: bool = False
    pdf_text_with_map: bool = False
    custody_markdown_bundle: bool = False


def load_config(path: Path | None = None) -> Config:
    config_path = path or DEFAULT_CONFIG
    raw = json.loads(config_path.read_text(encoding="utf-8"))
    mineru_backend = str(raw.get("mineru_backend", "pipeline"))
    mineru_parse_method = str(raw.get("mineru_parse_method", "auto"))
    mineru_language = str(raw.get("mineru_language", "ch"))
    preserve_images = raw.get("preserve_webarchive_images", False)
    if type(preserve_images) is not bool:
        raise ValueError("preserve_webarchive_images must be boolean")
    pdf_text_with_map = raw.get("pdf_text_with_map", False)
    if type(pdf_text_with_map) is not bool:
        raise ValueError("pdf_text_with_map must be boolean")
    custody_bundle = raw.get("custody_markdown_bundle", False)
    if type(custody_bundle) is not bool:
        raise ValueError("custody_markdown_bundle must be boolean")
    if mineru_backend not in MINERU_BACKENDS:
        raise ValueError(f"不支持的本机 MinerU 后端: {mineru_backend}")
    if mineru_parse_method not in MINERU_PARSE_METHODS:
        raise ValueError(f"不支持的 MinerU 解析模式: {mineru_parse_method}")
    if mineru_language not in MINERU_LANGUAGES:
        raise ValueError(f"不支持的 MinerU 语言: {mineru_language}")
    return Config(
        icloud_inbox=Path(raw["icloud_inbox"]).expanduser(),
        local_transfer_station=Path(raw["local_transfer_station"]).expanduser(),
        conversion_tool_workspace=Path(raw["conversion_tool_workspace"]).expanduser(),
        result_prefix=str(raw.get("result_prefix", "成果")),
        default_topic=str(raw.get("default_topic", "知识库转换")),
        markdown_preview_rows_per_sheet=int(raw.get("markdown_preview_rows_per_sheet", 200)),
        skip_name_prefixes=tuple(raw.get("skip_name_prefixes", [])),
        skip_name_contains=tuple(raw.get("skip_name_contains", [])),
        mineru_backend=mineru_backend,
        mineru_parse_method=mineru_parse_method,
        mineru_language=mineru_language,
        preserve_webarchive_images=preserve_images,
        pdf_text_with_map=pdf_text_with_map,
        custody_markdown_bundle=custody_bundle,
    )


def safe_name(text: str) -> str:
    text = text.strip() or "未命名"
    text = re.sub(r"[\\/:*?\"<>|]+", "-", text)
    text = re.sub(r"\s+", "_", text)
    text = text.strip("._-")
    return text or "未命名"


def now_batch_id(topic: str) -> str:
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{stamp}_{safe_name(topic)}"


def should_skip_name(path: Path, cfg: Config) -> bool:
    name = path.name
    for prefix in cfg.skip_name_prefixes:
        if name.startswith(prefix):
            return True
    for fragment in cfg.skip_name_contains:
        if fragment and fragment in name:
            return True
    return False


def iter_candidate_files(inbox: Path, cfg: Config) -> list[Path]:
    candidates: list[Path] = []
    if not inbox.exists():
        return candidates

    def walk(base: Path) -> None:
        for child in sorted(base.iterdir(), key=lambda p: p.name):
            if should_skip_name(child, cfg):
                continue
            if child.is_dir():
                walk(child)
                continue
            if child.suffix.lower() == ".icloud":
                continue
            if child.suffix.lower() in SUPPORTED_EXTENSIONS:
                candidates.append(child)

    walk(inbox)
    return candidates


def select_candidates(cfg: Config, args: argparse.Namespace) -> list[Path]:
    candidates = iter_candidate_files(cfg.icloud_inbox, cfg)
    include = tuple(getattr(args, "include", None) or [])
    exclude = tuple(getattr(args, "exclude", None) or [])
    limit = getattr(args, "limit", None)

    if include:
        candidates = [
            path
            for path in candidates
            if any(fragment in path.relative_to(cfg.icloud_inbox).as_posix() for fragment in include)
        ]
    if exclude:
        candidates = [
            path
            for path in candidates
            if not any(fragment in path.relative_to(cfg.icloud_inbox).as_posix() for fragment in exclude)
        ]
    if limit is not None:
        candidates = candidates[: max(0, int(limit))]
    return candidates


def file_kind(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == '.zip':
        return 'image_zip'
    if ext in MINERU_EXTENSIONS:
        return "mineru"
    if ext in HEIC_EXTENSIONS:
        return "heic_to_mineru"
    if ext in DOCX_EXTENSIONS:
        return "docx_to_markdown"
    if ext in LEGACY_DOC_EXTENSIONS:
        return "doc_to_markdown"
    if ext in {".xlsx", ".xlsm"}:
        return "spreadsheet_to_markdown_csv"
    if ext == ".xls":
        return "xls_needs_conversion"
    if ext in TEXT_EXTENSIONS:
        return "text_direct"
    if ext == ".pptx":
        return "pptx_text_extract"
    if ext == ".ppt":
        return "ppt_needs_conversion"
    if ext in WEBARCHIVE_EXTENSIONS:
        return "webarchive_to_markdown"
    return "unsupported"


def stat_row(batch_id: str, original: Path, inbox: Path, staged: Path) -> dict[str, str]:
    stat = original.stat()
    rel = original.relative_to(inbox)
    kind = file_kind(original)
    action = "process" if kind != "unsupported" else "skip"
    return {
        "batch_id": batch_id,
        "original_path": str(original),
        "relative_input_path": rel.as_posix(),
        "group_id": "",
        "group_role": "",
        "group_main": "",
        "linked_from": "",
        "link_target": "",
        "staged_path": str(staged),
        "file_name": original.name,
        "extension": original.suffix.lower(),
        "size_bytes": str(stat.st_size),
        "modified_time": dt.datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
        "file_kind": kind,
        "action": action,
        "processing_status": "copied" if action == "process" else "skipped",
        "kb_folder": "",
        "classification": "",
        "primary_input": "",
        "auxiliary_input": "",
        "quality_status": "",
        "source_url": "",
        "conversion_method": "",
        "conversion_status": "",
        "notes": "",
    }


def read_manifest(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_manifest(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=MANIFEST_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in MANIFEST_FIELDS})


def batch_dirs(batch: Path) -> dict[str, Path]:
    return {
        "raw": batch / "00_原始输入",
        "pre": batch / "01_预处理",
        "mineru": batch / "02_MinerU原始输出",
        "direct": batch / "03_直接转换输出",
        "kb": batch / "04_知识库输入文件",
        "logs": batch / "90_日志与清单",
        "review": batch / "99_需人工复核",
    }


def ensure_batch_dirs(batch: Path) -> None:
    for path in batch_dirs(batch).values():
        path.mkdir(parents=True, exist_ok=True)


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem = path.name[: -len(path.suffix)] if path.suffix else path.name
    suffix = path.suffix
    for i in range(2, 1000):
        candidate = path.with_name(f"{stem}_{i}{suffix}")
        if not candidate.exists():
            return candidate
    digest = hashlib.sha1(str(path).encode("utf-8")).hexdigest()[:8]
    return path.with_name(f"{stem}_{digest}{suffix}")


def rel_to_batch(path_text: str, batch: Path) -> str:
    if not path_text:
        return ""
    try:
        return Path(path_text).resolve().relative_to(batch.resolve()).as_posix()
    except ValueError:
        return path_text


def copy_tree_file(source: Path, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, dest)


def cmd_plan(args: argparse.Namespace) -> int:
    cfg = load_config(args.config)
    candidates = select_candidates(cfg, args)
    rows = [stat_row("PLAN", path, cfg.icloud_inbox, Path("")) for path in candidates]
    annotate_relationships(rows, cfg.icloud_inbox)
    missing_links = find_missing_markdown_links(rows, cfg.icloud_inbox)
    print(f"iCloud 投递箱: {cfg.icloud_inbox}")
    print(f"候选文件数: {len(candidates)}")
    by_kind: dict[str, int] = {}
    for row in rows:
        by_kind[row["file_kind"]] = by_kind.get(row["file_kind"], 0) + 1
    for kind, count in sorted(by_kind.items()):
        print(f"- {kind}: {count}")
    groups = {row["group_id"] for row in rows if row.get("group_id")}
    if groups:
        print(f"资料组数: {len(groups)}")
        for group_id in sorted(groups):
            group_rows = [row for row in rows if row.get("group_id") == group_id]
            print(f"- {group_id}: {len(group_rows)} 个关联文件")
    if missing_links:
        print(f"缺失附件链接: {len(missing_links)}")
        if args.verbose:
            for main_rel, target in missing_links:
                print(f"MISSING link_from={main_rel} target={target}")
    if args.verbose:
        for row in rows:
            group = row.get("group_id") or "-"
            role = row.get("group_role") or "-"
            print(f"{row['file_kind']:28s} group={group:28s} role={role:18s} {row['original_path']}")
    return 0


def cmd_prepare(args: argparse.Namespace) -> int:
    cfg = load_config(args.config)
    topic = args.topic or cfg.default_topic
    candidates = select_candidates(cfg, args)
    if not candidates:
        print("没有找到需要处理的候选文件。")
        return 0

    batch_id = now_batch_id(topic)
    batch = cfg.local_transfer_station / batch_id
    ensure_batch_dirs(batch)

    rows: list[dict[str, str]] = []
    for original in candidates:
        rel = original.relative_to(cfg.icloud_inbox)
        staged = batch_dirs(batch)["raw"] / rel
        staged = unique_path(staged)
        copy_tree_file(original, staged)
        rows.append(stat_row(batch_id, original, cfg.icloud_inbox, staged))

    annotate_relationships(rows, cfg.icloud_inbox)
    write_manifest(batch / "_batch_manifest.csv", rows)
    write_batch_readme(batch, topic, cfg, rows)
    print(f"已创建批次: {batch}")
    print(f"已复制候选文件: {len(rows)}")
    return 0


def env_timeout_seconds(name: str) -> int | None:
    raw = os.environ.get(name, "").strip()
    if not raw:
        return None
    try:
        value = int(raw)
    except ValueError:
        return None
    return value if value > 0 else None


def env_float(name: str, default: float) -> float:
    raw = os.environ.get(name, "").strip()
    if not raw:
        return default
    try:
        value = float(raw)
    except ValueError:
        return default
    return value if value > 0 else default


def run_command(command: list[str], cwd: Path | None = None, timeout: int | None = None) -> tuple[bool, str]:
    try:
        process = subprocess.Popen(
            command,
            cwd=str(cwd) if cwd else None,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            start_new_session=True,
        )
    except FileNotFoundError as exc:
        return False, str(exc)
    try:
        output, _ = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        try:
            os.killpg(process.pid, signal.SIGTERM)
        except ProcessLookupError:
            pass
        try:
            output, _ = process.communicate(timeout=10)
        except subprocess.TimeoutExpired:
            try:
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
            output, _ = process.communicate()
        timeout_note = f"命令超时 {timeout} 秒"
        output = f"{timeout_note}\n{output or ''}".strip()
        return False, output
    return process.returncode == 0, (output or "").strip()


def cmd_process(args: argparse.Namespace) -> int:
    cfg = load_config(args.config)
    batch = args.batch.resolve()
    ensure_batch_dirs(batch)
    manifest = batch / "_batch_manifest.csv"
    rows = read_manifest(manifest)

    processed = 0
    for row in rows:
        if row.get("action") != "process":
            continue
        status = row.get("processing_status", "")
        if status not in {"copied", "failed", ""} and not args.force:
            continue
        process_one(row, batch, cfg)
        processed += 1
    remove_junk_files(batch_dirs(batch)["kb"])
    write_relation_index(batch, rows)
    write_manifest(manifest, rows)
    write_kb_manifest(batch, rows)
    write_kb_readme(batch, rows)
    print(f"处理完成或已跳过: {processed} 条记录")
    print(f"知识库输入目录: {batch_dirs(batch)['kb']}")
    return 0


def process_one(row: dict[str, str], batch: Path, cfg: Config) -> None:
    source = Path(row["staged_path"])
    kind = row.get("file_kind") or file_kind(source)
    try:
        if kind == "mineru":
            process_mineru(row, source, batch, cfg)
        elif kind == "heic_to_mineru":
            process_heic(row, source, batch, cfg)
        elif kind == "docx_to_markdown":
            process_docx(row, source, batch, cfg)
        elif kind == "doc_to_markdown":
            process_legacy_doc(row, source, batch, cfg)
        elif kind == "spreadsheet_to_markdown_csv":
            process_xlsx(row, source, batch, cfg)
        elif kind == "xls_needs_conversion":
            mark_review(row, source, batch, "xls 需要先安装 xlrd 或手动另存为 xlsx 后处理")
        elif kind == "text_direct":
            process_text_direct(row, source, batch)
        elif kind == "pptx_text_extract":
            process_pptx(row, source, batch)
        elif kind == "ppt_needs_conversion":
            mark_review(row, source, batch, "ppt 需要先转为 pptx 或 pdf")
        elif kind == "webarchive_to_markdown":
            process_webarchive(row, source, batch, preserve_images=cfg.preserve_webarchive_images)
        elif kind == 'image_zip' and cfg.custody_markdown_bundle:
            from markdown_bundle import convert_markdown_bundle
            destination = item_output_dir(batch_dirs(batch)['kb'], source)
            primary = convert_markdown_bundle(source, destination)
            row.update(processing_status='processed', kb_folder=rel_to_batch(str(destination), batch),
                       classification='混合型', primary_input=rel_to_batch(str(primary), batch),
                       auxiliary_input='', quality_status='', conversion_method='markdown_bundle_offline',
                       conversion_status='processed', notes='单篇 Markdown 资料包直接导入；原图与辅助文件保留，未重复调用 OCR')
        elif kind == 'image_zip':
            from image_zip_converter import convert_image_zip
            import tempfile
            work = Path(tempfile.mkdtemp(prefix='image-zip-', dir=batch_dirs(batch)['pre'])) / 'work'
            destination = item_output_dir(batch_dirs(batch)['kb'], source)
            count = convert_image_zip(source, work_dir=work, output_dir=destination, run_command=run_command,
                                      tools_dir=cfg.conversion_tool_workspace, language=cfg.mineru_language,
                                      timeout=env_timeout_seconds('KB_WORKFLOW_MINERU_TIMEOUT_SECONDS'))
            order_report=json.loads((work/'image-order.json').read_text(encoding='utf-8'))
            edge_review=any(i.get('edge_review') for i in order_report['images'])
            row.update(processing_status='processed', kb_folder=rel_to_batch(str(destination),batch),
                       classification='文字型', primary_input=rel_to_batch(str(destination/'正文.md'),batch),
                       auxiliary_input=rel_to_batch(str(destination/'images'),batch), quality_status='needs_manual_review' if edge_review else '',
                       conversion_method='image_zip_local_ocr', conversion_status='processed',
                       notes=f'图片 ZIP 按文件名数字排序，共 {count} 张，合并为一份文章；原图与排序记录保留'+('；部分图片边缘可能截断文字，需对照原图抽查 OCR' if edge_review else ''))
        else:
            mark_review(row, source, batch, f"暂不支持的文件类型: {source.suffix}")
    except Exception as exc:  # noqa: BLE001
        row["processing_status"] = "failed"
        row["quality_status"] = "needs_manual_review"
        row["notes"] = f"{row.get('notes', '')} | {type(exc).__name__}: {exc}".strip(" |")


def annotate_relationships(rows: list[dict[str, str]], inbox: Path) -> None:
    by_original = {str(Path(row["original_path"]).resolve()): row for row in rows}
    main_md_rows = [
        row
        for row in rows
        if Path(row["original_path"]).suffix.lower() == ".md"
        and row.get("file_kind") == "text_direct"
    ]
    main_by_stem = {Path(row["original_path"]).stem: row for row in main_md_rows}

    for main in main_md_rows:
        main_path = Path(main["original_path"])
        group_id = safe_name(main_path.stem)
        main_rel = main_path.relative_to(inbox).as_posix()
        main.update(
            {
                "group_id": group_id,
                "group_role": "main_markdown",
                "group_main": main_rel,
                "linked_from": "",
                "link_target": "",
            }
        )
        for raw_target in extract_markdown_link_targets(main_path):
            target_path = resolve_markdown_target(main_path, raw_target)
            if target_path is None:
                continue
            target_row = by_original.get(str(target_path.resolve()))
            if not target_row:
                continue
            target_row.update(
                {
                    "group_id": group_id,
                    "group_role": "linked_attachment",
                    "group_main": main_rel,
                    "linked_from": main_rel,
                    "link_target": raw_target,
                }
            )

    for row in rows:
        if row.get("group_id"):
            continue
        source = Path(row["original_path"])
        main = main_by_stem.get(source.parent.name)
        if not main:
            continue
        main_path = Path(main["original_path"])
        row.update(
            {
                "group_id": safe_name(main_path.stem),
                "group_role": "sibling_attachment",
                "group_main": main_path.relative_to(inbox).as_posix(),
                "linked_from": main_path.relative_to(inbox).as_posix(),
                "link_target": source.relative_to(main_path.parent).as_posix()
                if source.is_relative_to(main_path.parent)
                else source.name,
            }
        )


def extract_markdown_link_targets(path: Path) -> list[str]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return []
    return [match.strip() for match in re.findall(r"!?\[[^\]]*\]\(([^)]+)\)", text)]


def resolve_markdown_target(markdown_path: Path, target: str) -> Path | None:
    target = target.strip()
    if not target:
        return None
    if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target):
        return None
    target = target.split("#", 1)[0].strip()
    target = target.strip("<>")
    target = urllib.parse.unquote(target)
    if not target:
        return None
    return (markdown_path.parent / target).resolve()


def find_missing_markdown_links(rows: list[dict[str, str]], inbox: Path) -> list[tuple[str, str]]:
    missing: list[tuple[str, str]] = []
    for row in rows:
        source = Path(row["original_path"])
        if source.suffix.lower() != ".md":
            continue
        try:
            main_rel = source.relative_to(inbox).as_posix()
        except ValueError:
            main_rel = source.name
        for raw_target in extract_markdown_link_targets(source):
            target_path = resolve_markdown_target(source, raw_target)
            if target_path is None:
                continue
            if not target_path.exists():
                missing.append((main_rel, urllib.parse.unquote(raw_target)))
    return missing


def item_output_dir(base: Path, source: Path) -> Path:
    return base / safe_name(source.name[: -len(source.suffix)] if source.suffix else source.name)


def process_mineru(row: dict[str, str], source: Path, batch: Path, cfg: Config) -> None:
    if source.suffix.lower() == ".pdf" and cfg.pdf_text_with_map:
        # An explicit custody route, not an implicit MinerU/OCR fallback.
        from pdf_text_map import extract_pdf_text
        text, mapping = extract_pdf_text(source)
        out_dir = item_output_dir(batch_dirs(batch)["direct"], source)
        out_dir.mkdir(parents=True, exist_ok=True)
        out_md = out_dir / f"{safe_name(source.stem)}.md"
        out_md.write_bytes(text.encode("utf-8"))
        out_md.with_suffix(".pages.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
        copy_dir_contents(out_dir, kb_dir)
        row.update(processing_status="processed", kb_folder=rel_to_batch(str(kb_dir), batch),
                   classification="文字型", primary_input=rel_to_batch(str(kb_dir / out_md.name), batch),
                   auxiliary_input="", quality_status="needs_manual_review",
                   conversion_method="pypdf_text_mapped", conversion_status="processed",
                   notes="本机 PDF 文本层提取；保留逐页映射。图像、版面、表格及文本完整性仍需复核。")
        return
    if source.suffix.lower() == ".pdf" and os.environ.get("KB_WORKFLOW_PDF_TEXT_ONLY", "").lower() in {
        "1",
        "true",
        "yes",
    }:
        process_pdf_text_fallback(
            row,
            source,
            batch,
            "按本轮策略跳过 MinerU，直接抽取 PDF 文本层",
            method="pypdf_text_direct",
            status="processed",
        )
        return

    dirs = batch_dirs(batch)
    out_dir = item_output_dir(dirs["mineru"], source)
    timeout = env_timeout_seconds("KB_WORKFLOW_MINERU_TIMEOUT_SECONDS")
    ok, output = run_command(
        [
            str(cfg.conversion_tool_workspace / "pdf2md.sh"),
            str(source),
            str(out_dir),
            cfg.mineru_backend,
            cfg.mineru_parse_method,
            cfg.mineru_language,
        ],
        timeout=timeout,
    )
    if not ok:
        if source.suffix.lower() == ".pdf":
            process_pdf_text_fallback(
                row,
                source,
                batch,
                compact_note("MinerU 失败或超时，已改用 PDF 文本层降级抽取", output),
            )
            return
        if source.suffix.lower() in MINERU_EXTENSIONS and source.suffix.lower() != ".pdf":
            process_image_fallback(row, source, batch, compact_note("MinerU 失败，保留原图作为降级输入", output))
            return
        row["processing_status"] = "failed"
        row["quality_status"] = "needs_manual_review"
        row["notes"] = compact_note("MinerU 失败", output)
        return
    collect_mineru_result(row, source, out_dir, batch)


def process_pdf_text_fallback(
    row: dict[str, str],
    source: Path,
    batch: Path,
    reason: str,
    method: str = "pypdf_text_fallback",
    status: str = "fallback",
) -> None:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        mark_review(row, source, batch, compact_note(reason, "当前 Python 环境缺少 pypdf，无法降级抽取文本"))
        return

    try:
        reader = PdfReader(str(source))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                pass
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            text = normalize_pdf_text(page.extract_text() or "")
            if text:
                pages.append((index, text))
    except Exception as exc:  # noqa: BLE001
        mark_review(row, source, batch, compact_note(reason, f"PDF 文本层抽取失败: {type(exc).__name__}: {exc}"))
        return

    total_chars = sum(len(text) for _, text in pages)
    if total_chars < 80:
        process_pdf_page_image_fallback(
            row,
            source,
            batch,
            compact_note(reason, f"PDF 文本层内容过短: {total_chars} 字符"),
        )
        return

    out_dir = item_output_dir(batch_dirs(batch)["direct"], source)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_md = out_dir / f"{safe_name(source.stem)}.md"
    lines = [
        f"# {source.stem}",
        "",
        f"来源文件：`{source.name}`",
        "",
        "> 本文件由知识库流程在 MinerU 失败或超时后，使用 PDF 文本层降级抽取生成。版面、脚注、公式和表格需人工抽查。",
        "",
    ]
    for page_number, text in pages:
        lines.extend([f"## Page {page_number}", "", text, ""])
    out_md.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")

    kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
    copy_dir_contents(out_dir, kb_dir)
    row.update(
        {
            "processing_status": "processed",
            "kb_folder": rel_to_batch(str(kb_dir), batch),
            "classification": "文字型",
            "primary_input": rel_to_batch(str(kb_dir / out_md.name), batch),
            "auxiliary_input": "",
            "quality_status": "needs_manual_review",
            "conversion_method": method,
            "conversion_status": status,
            "notes": compact_note(reason, f"抽取 {len(pages)} 页，约 {total_chars} 字符"),
        }
    )


def process_pdf_page_image_fallback(row: dict[str, str], source: Path, batch: Path, reason: str) -> None:
    try:
        import pypdfium2 as pdfium  # type: ignore
    except ImportError:
        mark_review(row, source, batch, compact_note(reason, "当前 Python 环境缺少 pypdfium2，无法渲染页图"))
        return

    out_dir = item_output_dir(batch_dirs(batch)["direct"], source)
    images_dir = out_dir / "images"
    if images_dir.exists():
        shutil.rmtree(images_dir)
    images_dir.mkdir(parents=True, exist_ok=True)

    try:
        pdf = pdfium.PdfDocument(str(source))
        scale = env_float("KB_WORKFLOW_PDF_IMAGE_SCALE", 6.0)
        image_paths: list[Path] = []
        for index in range(len(pdf)):
            page = pdf[index]
            try:
                image = page.render(scale=scale).to_pil().convert("RGB")
            finally:
                page.close()
            image_path = images_dir / f"page_{index + 1:03d}.jpg"
            image.save(image_path, quality=86, optimize=True)
            image_paths.append(image_path)
        pdf.close()
    except Exception as exc:  # noqa: BLE001
        mark_review(row, source, batch, compact_note(reason, f"PDF 页图渲染失败: {type(exc).__name__}: {exc}"))
        return

    if not image_paths:
        mark_review(row, source, batch, compact_note(reason, "PDF 页图渲染结果为空"))
        return

    out_md = out_dir / f"{safe_name(source.stem)}.md"
    lines = [
        f"# {source.stem}",
        "",
        f"来源文件：`{source.name}`",
        "",
        "> 本文件未检测到可抽取文本层，已渲染为逐页图片。可在 Obsidian 中直接阅读原始页图；全文检索需后续 OCR 校对。",
        "",
    ]
    for index, image_path in enumerate(image_paths, start=1):
        lines.extend([f"## Page {index}", "", f"![Page {index}](images/{image_path.name})", ""])
    out_md.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")

    kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
    copy_dir_contents(out_dir, kb_dir)
    row.update(
        {
            "processing_status": "processed",
            "kb_folder": rel_to_batch(str(kb_dir), batch),
            "classification": "图片型",
            "primary_input": rel_to_batch(str(kb_dir / out_md.name), batch),
            "auxiliary_input": rel_to_batch(str(kb_dir / "images"), batch),
            "quality_status": "needs_manual_review",
            "conversion_method": "pdfium_page_images",
            "conversion_status": "fallback",
            "notes": compact_note(reason, f"已渲染 {len(image_paths)} 页图片"),
        }
    )


def normalize_pdf_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    raw_lines = [re.sub(r"\s+", " ", line).strip() for line in normalized.splitlines()]
    paragraphs: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        paragraph = " ".join(buffer)
        paragraph = re.sub(r"([A-Za-z])-\s+([a-z])", r"\1\2", paragraph)
        paragraph = re.sub(r"\s+", " ", paragraph).strip()
        if paragraph:
            paragraphs.append(paragraph)
        buffer.clear()

    for line in raw_lines:
        if not line:
            flush()
            continue
        if looks_like_pdf_heading(line):
            flush()
            paragraphs.append(f"### {line}")
            continue
        buffer.append(line)
    flush()
    return "\n\n".join(paragraphs)


def looks_like_pdf_heading(line: str) -> bool:
    if len(line) > 120:
        return False
    if re.fullmatch(r"(abstract|introduction|references|conclusion|discussion|appendix)", line, re.I):
        return True
    if re.fullmatch(r"\d+(?:\.\d+)*\.?\s+[A-Z][A-Za-z0-9,;:()'’\-\s]{2,}", line):
        return True
    words = re.findall(r"[A-Za-z]+", line)
    return bool(words) and len(words) <= 12 and sum(word[:1].isupper() for word in words) >= max(2, len(words) - 1)


def process_image_fallback(row: dict[str, str], source: Path, batch: Path, note: str) -> None:
    kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
    kb_dir.mkdir(parents=True, exist_ok=True)
    dest = kb_dir / source.name
    shutil.copy2(source, dest)
    row.update(
        {
            "processing_status": "needs_manual_review",
            "kb_folder": rel_to_batch(str(kb_dir), batch),
            "classification": "图片型",
            "primary_input": rel_to_batch(str(dest), batch),
            "auxiliary_input": "",
            "quality_status": "needs_manual_review",
            "notes": note,
        }
    )


def process_heic(row: dict[str, str], source: Path, batch: Path, cfg: Config) -> None:
    pre_dir = item_output_dir(batch_dirs(batch)["pre"], source)
    pre_dir.mkdir(parents=True, exist_ok=True)
    converted = pre_dir / f"{safe_name(source.stem)}.jpg"
    ok, output = run_command(["sips", "-s", "format", "jpeg", str(source), "--out", str(converted)])
    if not ok:
        row["processing_status"] = "failed"
        row["quality_status"] = "needs_manual_review"
        row["notes"] = compact_note("HEIC 转 JPEG 失败", output)
        return
    process_mineru(row, converted, batch, cfg)


def process_docx(row: dict[str, str], source: Path, batch: Path, cfg: Config) -> None:
    out_dir = item_output_dir(batch_dirs(batch)["direct"], source)
    out_md = out_dir / f"{safe_name(source.stem)}.md"
    tool = cfg.conversion_tool_workspace / "convert_docx_direct_to_md.py"
    ok, output = run_command([sys.executable, str(tool), str(source), str(out_md)])
    if not ok:
        row["processing_status"] = "failed"
        row["quality_status"] = "needs_manual_review"
        row["notes"] = compact_note("DOCX 转 Markdown 失败", output)
        return
    clean_word_field_artifacts(out_md)
    media_count, appended_count = ensure_docx_media_referenced(source, out_md)
    kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
    copy_dir_contents(out_dir, kb_dir)
    note = "DOCX 直接转 Markdown"
    if media_count:
        note += f"，保留 {media_count} 张内嵌图片"
    if appended_count:
        note += f"（其中 {appended_count} 张通过媒体兜底索引补充）"
    row.update(
        {
            "processing_status": "processed",
            "kb_folder": rel_to_batch(str(kb_dir), batch),
            "classification": "文字型",
            "primary_input": rel_to_batch(str(kb_dir / out_md.name), batch),
            "auxiliary_input": images_if_present(kb_dir, batch),
            "quality_status": "",
            "notes": note,
        }
    )


def clean_word_field_artifacts(path: Path) -> None:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    cleaned: list[str] = []
    for line in lines:
        stripped = line.strip()
        if re.search(r"\b(?:TOC \\o|HYPERLINK \\l|PAGEREF _Toc)\b", stripped):
            continue
        if re.fullmatch(r"PAGE(?:\s+\\\*\s+MERGEFORMAT)?\s*\d*", stripped):
            continue
        cleaned.append(line)
    normalized = re.sub(r"\n{3,}", "\n\n", "\n".join(cleaned))
    path.write_text(normalized.rstrip() + "\n", encoding="utf-8")


def ensure_docx_media_referenced(source: Path, output_md: Path) -> tuple[int, int]:
    media_dir = output_md.parent / "images"
    text = output_md.read_text(encoding="utf-8", errors="replace")
    referenced = set(re.findall(r"\]\(images/([^)]+)\)", text))
    appended: list[str] = []

    with zipfile.ZipFile(source) as zf:
        media_names = [
            name
            for name in zf.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        ]
        media_names.sort(key=lambda name: natural_sort_key(Path(name).name))
        for archive_name in media_names:
            file_name = Path(archive_name).name
            dest = media_dir / file_name
            if not dest.exists():
                media_dir.mkdir(parents=True, exist_ok=True)
                with zf.open(archive_name) as src, dest.open("wb") as out:
                    shutil.copyfileobj(src, out)
            if file_name not in referenced:
                appended.append(file_name)

    if appended:
        lines = [text.rstrip(), "", "## 文档内嵌图片", ""]
        for index, file_name in enumerate(appended, start=1):
            lines.append(f"![文档内嵌图片 {index}](images/{file_name})")
            lines.append("")
        output_md.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")

    return len(media_names), len(appended)


def natural_sort_key(value: str) -> list[object]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value)]


def process_legacy_doc(row: dict[str, str], source: Path, batch: Path, cfg: Config) -> None:
    pre_dir = item_output_dir(batch_dirs(batch)["pre"], source)
    pre_dir.mkdir(parents=True, exist_ok=True)
    converted = pre_dir / f"{safe_name(source.stem)}.docx"
    ok, output = run_command(
        ["textutil", "-convert", "docx", "-output", str(converted), str(source)]
    )
    if not ok or not converted.exists():
        row["processing_status"] = "failed"
        row["quality_status"] = "needs_manual_review"
        row["notes"] = compact_note("DOC 转 DOCX 失败", output)
        return
    process_docx(row, converted, batch, cfg)
    if row.get("processing_status") == "processed":
        row["notes"] = "DOC 经 textutil 转 DOCX 后转 Markdown"


def process_xlsx(row: dict[str, str], source: Path, batch: Path, cfg: Config) -> None:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        mark_review(row, source, batch, "当前 Python 环境缺少 openpyxl，无法直接处理 xlsx")
        return

    out_dir = item_output_dir(batch_dirs(batch)["direct"], source)
    out_dir.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    md_lines = [f"# {source.stem}", "", f"来源文件：`{source.name}`", ""]
    csv_paths: list[Path] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        safe_sheet = safe_name(sheet.title)
        csv_path = out_dir / f"{safe_sheet}.csv"
        csv_paths.append(csv_path)
        with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            for values in rows:
                writer.writerow(["" if value is None else value for value in values])
        md_lines.append(f"## {sheet.title}")
        md_lines.append("")
        md_lines.append(f"- 行数：{len(rows)}")
        md_lines.append(f"- 列数：{max((len(r) for r in rows), default=0)}")
        md_lines.append(f"- 完整 CSV：`{csv_path.name}`")
        md_lines.append("")
        preview_rows = rows[: cfg.markdown_preview_rows_per_sheet]
        if preview_rows:
            md_lines.extend(markdown_table(preview_rows))
            md_lines.append("")
        if len(rows) > len(preview_rows):
            md_lines.append(f"> Markdown 仅展示前 {len(preview_rows)} 行，完整数据见 CSV。")
            md_lines.append("")
    out_md = out_dir / f"{safe_name(source.stem)}.md"
    out_md.write_text("\n".join(md_lines).rstrip() + "\n", encoding="utf-8")

    kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
    copy_dir_contents(out_dir, kb_dir)
    row.update(
        {
            "processing_status": "processed",
            "kb_folder": rel_to_batch(str(kb_dir), batch),
            "classification": "表格型",
            "primary_input": rel_to_batch(str(kb_dir / out_md.name), batch),
            "auxiliary_input": append_paths("", [rel_to_batch(str(kb_dir / p.name), batch) for p in csv_paths]),
            "quality_status": "",
            "notes": "XLSX 转 Markdown + CSV",
        }
    )


def process_text_direct(row: dict[str, str], source: Path, batch: Path) -> None:
    kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
    kb_dir.mkdir(parents=True, exist_ok=True)
    if source.suffix.lower() == ".txt":
        content = source.read_text(encoding="utf-8", errors="replace").strip()
        out = kb_dir / f"{safe_name(source.stem)}.md"
        out.write_text(f"# {source.stem}\n\n{content}\n", encoding="utf-8")
        primary = out
        note = "TXT 转 Markdown"
    else:
        out = kb_dir / source.name
        shutil.copy2(source, out)
        if out.suffix.lower() == ".md":
            normalized = extract_data_uri_images(
                out.read_text(encoding="utf-8", errors="replace"),
                out.parent,
                f"{out.stem}_内嵌图片",
            )
            out.write_text(normalized, encoding="utf-8")
        primary = out
        note = "原文本直接纳入"
    classification = "表格型" if source.suffix.lower() == ".csv" else "文字型"
    row.update(
        {
            "processing_status": "processed",
            "kb_folder": rel_to_batch(str(kb_dir), batch),
            "classification": classification,
            "primary_input": rel_to_batch(str(primary), batch),
            "auxiliary_input": "",
            "quality_status": "",
            "notes": note,
        }
    )


def process_pptx(row: dict[str, str], source: Path, batch: Path) -> None:
    out_dir = item_output_dir(batch_dirs(batch)["direct"], source)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_md = out_dir / f"{safe_name(source.stem)}.md"
    slides = extract_pptx_text(source)
    lines = [f"# {source.stem}", "", f"来源文件：`{source.name}`", ""]
    for i, texts in enumerate(slides, start=1):
        lines.append(f"## Slide {i}")
        lines.append("")
        if texts:
            for text in texts:
                lines.append(text)
                lines.append("")
        else:
            lines.append("> 未抽取到可读文本。")
            lines.append("")
    out_md.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
    copy_dir_contents(out_dir, kb_dir)
    row.update(
        {
            "processing_status": "processed",
            "kb_folder": rel_to_batch(str(kb_dir), batch),
            "classification": "混合型",
            "primary_input": rel_to_batch(str(kb_dir / out_md.name), batch),
            "auxiliary_input": "",
            "quality_status": "needs_manual_review",
            "notes": "PPTX 仅抽取文本；复杂版面建议转 PDF 后用 MinerU 核对",
        }
    )


def process_webarchive(row: dict[str, str], source: Path, batch: Path, *, preserve_images: bool = False) -> None:
    out_dir = item_output_dir(batch_dirs(batch)["direct"], source)
    result = convert_webarchive_to_markdown(source, out_dir, preserve_images=preserve_images)
    metadata = result.metadata
    common = {
        "source_url": str(metadata.get("source_url", "")),
        "conversion_method": str(metadata.get("conversion_method", "")),
        "conversion_status": str(metadata.get("conversion_status", result.status)),
    }
    if result.status != "processed" or result.markdown_path is None:
        review_dir = item_output_dir(batch_dirs(batch)["review"], source)
        review_dir.mkdir(parents=True, exist_ok=True)
        source_dest = review_dir / source.name
        shutil.copy2(source, source_dest)
        metadata_dest = review_dir / result.metadata_path.name
        shutil.copy2(result.metadata_path, metadata_dest)
        row.update(
            {
                **common,
                "processing_status": "needs_manual_review",
                "kb_folder": rel_to_batch(str(review_dir), batch),
                "classification": "待判定",
                "primary_input": "",
                "auxiliary_input": append_paths(
                    "",
                    [rel_to_batch(str(source_dest), batch), rel_to_batch(str(metadata_dest), batch)],
                ),
                "quality_status": "needs_manual_review",
                "notes": result.notes,
            }
        )
        return

    kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
    copy_dir_contents(out_dir, kb_dir)
    kb_md = kb_dir / result.markdown_path.name
    kb_metadata = kb_dir / result.metadata_path.name
    quality_status = "needs_manual_review" if metadata.get("conversion_warnings") else ""
    row.update(
        {
            **common,
            "processing_status": "processed",
            "kb_folder": rel_to_batch(str(kb_dir), batch),
            "classification": "文字型",
            "primary_input": rel_to_batch(str(kb_md), batch),
            "auxiliary_input": rel_to_batch(str(kb_metadata), batch),
            "quality_status": quality_status,
            "notes": result.notes,
        }
    )


def extract_pptx_text(source: Path) -> list[list[str]]:
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    slides: list[list[str]] = []
    with zipfile.ZipFile(source) as zf:
        names = sorted(
            [name for name in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", name)],
            key=lambda n: int(re.search(r"slide(\d+)\.xml$", n).group(1)),  # type: ignore[union-attr]
        )
        for name in names:
            root = ET.fromstring(zf.read(name))
            texts: list[str] = []
            for paragraph in root.findall(".//a:p", ns):
                pieces = [node.text or "" for node in paragraph.findall(".//a:t", ns)]
                text = " ".join(piece.strip() for piece in pieces if piece.strip()).strip()
                if text:
                    texts.append(text)
            slides.append(texts)
    return slides


def collect_mineru_result(row: dict[str, str], source: Path, out_dir: Path, batch: Path) -> None:
    md_files = sorted(out_dir.rglob("*.md"))
    json_files = sorted(out_dir.rglob("*content_list_v2.json"))
    image_dirs = sorted([p for p in out_dir.rglob("images") if p.is_dir()])
    kb_dir = item_output_dir(batch_dirs(batch)["kb"], source)
    kb_dir.mkdir(parents=True, exist_ok=True)

    selected: list[Path] = []
    md_text_len = 0
    if md_files:
        md = md_files[0]
        md_text_len = len(md.read_text(encoding="utf-8", errors="replace").strip())
        shutil.copy2(md, kb_dir / md.name)
        selected.append(kb_dir / md.name)
    if json_files:
        json_path = json_files[0]
        json_dest = kb_dir / "content_list_v2.json"
        shutil.copy2(json_path, json_dest)
        selected.append(json_dest)
    for image_dir in image_dirs[:1]:
        dest = kb_dir / "images"
        if dest.exists():
            shutil.rmtree(dest)
        shutil.copytree(image_dir, dest)

    image_count = count_files(kb_dir / "images")
    classification, primary = classify_mineru_output(source, md_files, json_files, md_text_len, image_count, kb_dir)
    if source.suffix.lower() in MINERU_EXTENSIONS and source.suffix.lower() != ".pdf":
        original_dest = kb_dir / source.name
        shutil.copy2(source, original_dest)
    aux_paths = []
    if (kb_dir / "images").exists():
        aux_paths.append(rel_to_batch(str(kb_dir / "images"), batch))
    if (kb_dir / source.name).exists():
        aux_paths.append(rel_to_batch(str(kb_dir / source.name), batch))

    if not primary:
        mark_review(row, source, batch, "MinerU 输出中未找到可用 md 或 content_list_v2.json")
        return

    note = f"MinerU 输出；md 字符数 {md_text_len}；图片数 {image_count}"
    quality = "needs_manual_review" if md_text_len < 80 and image_count > 0 else ""
    row.update(
        {
            "processing_status": "processed",
            "kb_folder": rel_to_batch(str(kb_dir), batch),
            "classification": classification,
            "primary_input": append_paths("", [rel_to_batch(str(p), batch) for p in primary]),
            "auxiliary_input": append_paths("", aux_paths),
            "quality_status": quality,
            "notes": note,
        }
    )


def classify_mineru_output(
    source: Path,
    md_files: list[Path],
    json_files: list[Path],
    md_text_len: int,
    image_count: int,
    kb_dir: Path,
) -> tuple[str, list[Path]]:
    md = kb_dir / md_files[0].name if md_files else None
    json_path = kb_dir / "content_list_v2.json" if json_files else None
    ext = source.suffix.lower()
    if json_path and (ext != ".pdf" or md_text_len < 500 or image_count >= 3):
        if md and md_text_len >= 80:
            return "混合型", [md, json_path]
        return "框图型", [json_path]
    if md:
        return "文字型", [md]
    if json_path:
        return "框图型", [json_path]
    return "未知", []


def mark_review(row: dict[str, str], source: Path, batch: Path, reason: str) -> None:
    review_dir = item_output_dir(batch_dirs(batch)["review"], source)
    review_dir.mkdir(parents=True, exist_ok=True)
    dest = review_dir / source.name
    if source.exists() and source.is_file():
        shutil.copy2(source, dest)
    row.update(
        {
            "processing_status": "needs_manual_review",
            "kb_folder": rel_to_batch(str(review_dir), batch),
            "classification": "待判定",
            "primary_input": "",
            "auxiliary_input": rel_to_batch(str(dest), batch) if dest.exists() else "",
            "quality_status": "needs_manual_review",
            "notes": reason,
        }
    )


def markdown_table(rows: list[tuple[object, ...]]) -> list[str]:
    width = max((len(row) for row in rows), default=0)
    if width == 0:
        return []
    normalized = [list(row) + [""] * (width - len(row)) for row in rows]
    header = [format_cell(value) for value in normalized[0]]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    for row in normalized[1:]:
        lines.append("| " + " | ".join(format_cell(value) for value in row) + " |")
    return lines


def format_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", "<br>").replace("|", r"\|")
    return text.strip()


def copy_dir_contents(source: Path, dest: Path) -> None:
    dest.mkdir(parents=True, exist_ok=True)
    for child in source.iterdir():
        if child.name in JUNK_FILE_NAMES:
            continue
        target = dest / child.name
        if child.is_dir():
            if target.exists():
                shutil.rmtree(target)
            shutil.copytree(child, target)
        else:
            shutil.copy2(child, target)


def count_files(path: Path) -> int:
    if not path.exists():
        return 0
    if path.is_file():
        return 1
    return sum(1 for p in path.rglob("*") if p.is_file())


def images_if_present(kb_dir: Path, batch: Path) -> str:
    image_dir = kb_dir / "images"
    if image_dir.exists() and count_files(image_dir):
        return rel_to_batch(str(image_dir), batch)
    return ""


def append_paths(existing: str, paths: Iterable[str]) -> str:
    values = [p for p in existing.split(";") if p]
    values.extend(p for p in paths if p)
    seen: set[str] = set()
    deduped: list[str] = []
    for value in values:
        if value not in seen:
            seen.add(value)
            deduped.append(value)
    return ";".join(deduped)


def compact_note(prefix: str, output: str) -> str:
    output = re.sub(r"\s+", " ", output).strip()
    if len(output) > 500:
        output = output[:500] + "..."
    return f"{prefix}: {output}" if output else prefix


def write_batch_readme(batch: Path, topic: str, cfg: Config, rows: list[dict[str, str]]) -> None:
    lines = [
        f"# {batch.name}",
        "",
        f"- 主题：{topic}",
        f"- 创建时间：{dt.datetime.now().isoformat(timespec='seconds')}",
        f"- iCloud 投递箱：`{cfg.icloud_inbox}`",
        f"- 本地中转站：`{cfg.local_transfer_station}`",
        f"- 文件数：{len(rows)}",
        "",
        "## 目录",
        "",
        "- `00_原始输入/`：从 iCloud 复制来的原始材料。",
        "- `01_预处理/`：HEIC 转图等预处理结果。",
        "- `02_MinerU原始输出/`：MinerU 完整输出。",
        "- `03_直接转换输出/`：DOCX、XLSX、PPTX 等直接转换结果。",
        "- `04_知识库输入文件/`：最终建议入库文件。",
        "- `90_日志与清单/`：质量报告和清单。",
        "- `99_需人工复核/`：失败或暂不适合自动入库材料。",
        "",
    ]
    (batch / "README.md").write_text("\n".join(lines), encoding="utf-8")


def write_kb_manifest(batch: Path, rows: list[dict[str, str]]) -> None:
    kb_dir = batch_dirs(batch)["kb"]
    kb_dir.mkdir(parents=True, exist_ok=True)
    path = kb_dir / "知识库输入文件清单.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=KB_MANIFEST_FIELDS)
        writer.writeheader()
        for row in rows:
            selected = "yes" if row.get("primary_input") else "no"
            writer.writerow(
                {
                    "source_file": row.get("original_path", ""),
                    "group_id": row.get("group_id", ""),
                    "group_role": row.get("group_role", ""),
                    "linked_from": row.get("linked_from", ""),
                    "classification": row.get("classification", ""),
                    "selected_for_kb": selected,
                    "primary_input": row.get("primary_input", ""),
                    "auxiliary_input": row.get("auxiliary_input", ""),
                    "needs_manual_review": "yes"
                    if row.get("quality_status") == "needs_manual_review"
                    or row.get("processing_status") == "needs_manual_review"
                    else "no",
                    "source_url": row.get("source_url", ""),
                    "conversion_method": row.get("conversion_method", ""),
                    "conversion_status": row.get("conversion_status", ""),
                    "notes": row.get("notes", ""),
                }
            )


def write_kb_readme(batch: Path, rows: list[dict[str, str]]) -> None:
    kb_dir = batch_dirs(batch)["kb"]
    total = len(rows)
    selected = sum(1 for row in rows if row.get("primary_input"))
    review = sum(
        1
        for row in rows
        if row.get("quality_status") == "needs_manual_review"
        or row.get("processing_status") == "needs_manual_review"
    )
    by_class: dict[str, int] = {}
    for row in rows:
        cls = row.get("classification") or "未分类"
        by_class[cls] = by_class.get(cls, 0) + 1
    lines = [
        "# 知识库输入文件",
        "",
        f"- 批次：`{batch.name}`",
        f"- 总处理记录：{total}",
        f"- 已选入库：{selected}",
        f"- 需人工复核：{review}",
        "",
        "## 类型统计",
        "",
    ]
    for cls, count in sorted(by_class.items()):
        lines.append(f"- {cls}: {count}")
    lines.extend(
        [
            "",
            "## 资料关系",
            "",
            "如本批次包含主 Markdown 与附件，详见 `资料关系索引.md` 和 `资料关系索引.csv`。",
            "",
            "## 清单",
            "",
            "详见 `知识库输入文件清单.csv`。",
            "",
        ]
    )
    (kb_dir / "README.md").write_text("\n".join(lines), encoding="utf-8")


def write_relation_index(batch: Path, rows: list[dict[str, str]]) -> None:
    kb_dir = batch_dirs(batch)["kb"]
    grouped: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        group_id = row.get("group_id", "")
        if group_id:
            grouped.setdefault(group_id, []).append(row)

    write_group_packages(kb_dir, grouped)

    csv_path = kb_dir / "资料关系索引.csv"
    csv_fields = [
        "group_id",
        "group_role",
        "group_main",
        "source_file",
        "linked_from",
        "link_target",
        "primary_input",
        "auxiliary_input",
        "classification",
        "processing_status",
        "quality_status",
        "notes",
    ]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=csv_fields)
        writer.writeheader()
        for row in rows:
            if not row.get("group_id"):
                continue
            writer.writerow({field: row.get(field, "") for field in csv_fields})

    lines = ["# 资料关系索引", ""]
    if not grouped:
        lines.append("本批次未识别到主 Markdown 与附件关系。")
    for group_id, group_rows in sorted(grouped.items()):
        lines.extend([f"## {group_id}", ""])
        main_rows = [row for row in group_rows if row.get("group_role") == "main_markdown"]
        attachment_rows = [row for row in group_rows if row.get("group_role") != "main_markdown"]
        if main_rows:
            main = main_rows[0]
            lines.append(f"- 主 Markdown 来源：`{main.get('original_path', '')}`")
            if main.get("primary_input"):
                lines.append(f"- 主 Markdown 入库文件：`{main.get('primary_input', '')}`")
        else:
            lines.append("- 主 Markdown：本次处理范围内未包含。")
        lines.append("")
        if attachment_rows:
            lines.append("| 角色 | 来源文件 | 原始链接 | 入库主输入 | 辅助文件 |")
            lines.append("|---|---|---|---|---|")
            for row in attachment_rows:
                lines.append(
                    "| "
                    + " | ".join(
                        [
                            format_cell(row.get("group_role", "")),
                            format_cell(row.get("original_path", "")),
                            format_cell(row.get("link_target", "")),
                            format_cell(row.get("primary_input", "")),
                            format_cell(row.get("auxiliary_input", "")),
                        ]
                    )
                    + " |"
                )
        else:
            lines.append("- 附件：无。")
        lines.append("")
    (kb_dir / "资料关系索引.md").write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_group_packages(kb_dir: Path, grouped: dict[str, list[dict[str, str]]]) -> None:
    for group_id, group_rows in grouped.items():
        main_rows = [row for row in group_rows if row.get("group_role") == "main_markdown"]
        if not main_rows:
            continue
        main = main_rows[0]
        source = Path(main.get("original_path", ""))
        try:
            source_text = source.read_text(encoding="utf-8", errors="replace")
        except OSError:
            continue

        by_link_target = {
            row.get("link_target", ""): row
            for row in group_rows
            if row.get("link_target") and row.get("group_role") != "main_markdown"
        }

        def replace_link(match: re.Match[str]) -> str:
            bang, label, target = match.groups()
            row = by_link_target.get(target.strip())
            if not row:
                return match.group(0)
            replacement = best_kb_link(row, want_image=bool(bang))
            if not replacement:
                return match.group(0)
            return f"{bang}[{label}]({replacement})"

        package_path = kb_dir / f"{safe_name(group_id)}_资料包.md"
        rewritten = re.sub(r"(!?)\[([^\]]*)\]\(([^)]+)\)", replace_link, source_text)
        rewritten = extract_data_uri_images(
            rewritten,
            package_path.parent,
            f"{package_path.stem}_内嵌图片",
        )
        package_lines = [
            f"# {group_id} 资料包",
            "",
            "> 本文件由主 Markdown 生成，并将已识别附件链接改写到本批次知识库输入文件。原始关系详见 `资料关系索引.md`。",
            "",
            rewritten.strip(),
            "",
            "## 附件转换结果",
            "",
            "| 附件 | 转换后主输入 | 辅助文件 |",
            "|---|---|---|",
        ]
        for row in group_rows:
            if row.get("group_role") == "main_markdown":
                continue
            package_lines.append(
                "| "
                + " | ".join(
                    [
                        format_cell(Path(row.get("original_path", "")).name),
                        format_cell(kb_relative_paths(row.get("primary_input", ""))),
                        format_cell(kb_relative_paths(row.get("auxiliary_input", ""))),
                    ]
                )
                + " |"
            )
        package_path.write_text("\n".join(package_lines).rstrip() + "\n", encoding="utf-8")
        package_rel = f"04_知识库输入文件/{package_path.name}"
        original_primary = main.get("primary_input", "")
        if original_primary and original_primary != package_rel:
            main["auxiliary_input"] = append_paths(main.get("auxiliary_input", ""), [original_primary])
        main["primary_input"] = package_rel
        main["kb_folder"] = "04_知识库输入文件"
        main["notes"] = "主 Markdown 资料包，附件链接已改写；原始 Markdown 作为辅助溯源"


def extract_data_uri_images(text: str, base_dir: Path, image_dir_name: str) -> str:
    pattern = re.compile(
        r"!\[([^\]]*)\]\(data:image/([^;,\s)]+);base64,([A-Za-z0-9+/=\s]+?)\)",
        re.S,
    )
    if "data:image/" not in text:
        return text
    image_dir = base_dir / safe_name(image_dir_name)
    counter = 0

    def replacement(match: re.Match[str]) -> str:
        nonlocal counter
        alt, mime, payload = match.groups()
        raw = re.sub(r"\s+", "", payload)
        try:
            data = base64.b64decode(raw, validate=False)
        except Exception:
            return match.group(0)
        counter += 1
        ext = image_extension(data, mime)
        image_dir.mkdir(parents=True, exist_ok=True)
        dest = image_dir / f"embedded_{counter:02d}.{ext}"
        dest.write_bytes(data)
        rel = dest.relative_to(base_dir).as_posix()
        return f"![{alt}]({rel})"

    return pattern.sub(replacement, text)


def image_extension(data: bytes, mime: str) -> str:
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if data.startswith(b"\xff\xd8\xff"):
        return "jpg"
    normalized = mime.lower()
    if normalized == "jpeg":
        return "jpg"
    return re.sub(r"[^a-z0-9]+", "", normalized) or "bin"


def best_kb_link(row: dict[str, str], want_image: bool) -> str:
    image_exts = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".tif", ".tiff"}
    aux = split_paths(row.get("auxiliary_input", ""))
    primary = split_paths(row.get("primary_input", ""))
    if want_image:
        for path in aux + primary:
            rel = strip_kb_prefix(path)
            if Path(rel).suffix.lower() in image_exts:
                return rel
    for path in primary:
        rel = strip_kb_prefix(path)
        if Path(rel).suffix.lower() == ".md":
            return rel
    for path in primary + aux:
        rel = strip_kb_prefix(path)
        if rel:
            return rel
    return ""


def kb_relative_paths(value: str) -> str:
    return ";".join(strip_kb_prefix(path) for path in split_paths(value))


def strip_kb_prefix(path: str) -> str:
    prefix = "04_知识库输入文件/"
    return path[len(prefix) :] if path.startswith(prefix) else path


def cmd_check(args: argparse.Namespace) -> int:
    batch = args.batch.resolve()
    rows = read_manifest(batch / "_batch_manifest.csv")
    issues: list[str] = []
    warnings: list[str] = []
    kb_dir = batch_dirs(batch)["kb"]
    remove_junk_files(kb_dir)
    write_relation_index(batch, rows)
    if not kb_dir.exists():
        issues.append("缺少 04_知识库输入文件 目录")
    if not (kb_dir / "README.md").exists():
        issues.append("缺少 README.md")
    if not (kb_dir / "知识库输入文件清单.csv").exists():
        issues.append("缺少 知识库输入文件清单.csv")

    for bad in PROCESS_FILE_NAMES:
        if list(kb_dir.rglob(bad)):
            issues.append(f"过程文件混入最终目录: {bad}")

    for row in rows:
        check_manifest_row(batch, row, issues, warnings)

    for row in rows:
        if row.get("primary_input"):
            if row.get("quality_status") != "needs_manual_review":
                row["quality_status"] = "pass"
        if any((row.get("file_name", "") in issue) for issue in issues + warnings):
            row["quality_status"] = "needs_manual_review"

    report = render_quality_report(batch, rows, issues, warnings)
    report_path = batch_dirs(batch)["logs"] / "quality_report.md"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(report, encoding="utf-8")

    write_manifest(batch / "_batch_manifest.csv", rows)
    write_relation_index(batch, rows)
    write_kb_manifest(batch, rows)
    write_kb_readme(batch, rows)

    print(f"质量报告: {report_path}")
    print(f"严重问题: {len(issues)}")
    print(f"提醒: {len(warnings)}")
    return 1 if issues else 0


def remove_junk_files(root: Path) -> None:
    if not root.exists():
        return
    for name in JUNK_FILE_NAMES:
        for path in root.rglob(name):
            if path.is_file():
                path.unlink()


def check_manifest_row(batch: Path, row: dict[str, str], issues: list[str], warnings: list[str]) -> None:
    name = row.get("file_name", "")
    if row.get("processing_status") in {"failed", "needs_manual_review"}:
        warnings.append(f"{name}: {row.get('processing_status')} - {row.get('notes', '')}")
    primary = split_paths(row.get("primary_input", ""))
    if row.get("processing_status") == "processed" and not primary:
        issues.append(f"{name}: 已处理但缺少 primary_input")
    for rel in primary:
        path = batch / rel
        if not path.exists():
            issues.append(f"{name}: 主输入缺失 {rel}")
            continue
        if path.is_file() and path.stat().st_size == 0:
            issues.append(f"{name}: 主输入为空 {rel}")
        if path.suffix.lower() == ".md":
            check_markdown(name, path, issues, warnings, structured_images=row.get("conversion_method") == "markdown_bundle_offline")
        elif path.name == "content_list_v2.json":
            check_json(name, path, issues)
        elif path.suffix.lower() == ".csv":
            check_csv(name, path, issues, warnings)
    for rel in split_paths(row.get("auxiliary_input", "")):
        path = batch / rel
        if not path.exists():
            warnings.append(f"{name}: 辅助输入缺失 {rel}")
        elif path.is_file() and path.stat().st_size == 0:
            warnings.append(f"{name}: 辅助输入为空 {rel}")
        elif path.suffix.lower() == ".json":
            check_json(name, path, issues)


def split_paths(value: str) -> list[str]:
    return [p for p in value.split(";") if p]


def check_markdown(name: str, path: Path, issues: list[str], warnings: list[str], *, structured_images: bool = False) -> None:
    text = path.read_text(encoding="utf-8", errors="replace")
    stripped = text.strip()
    if len(stripped) < 40:
        warnings.append(f"{name}: Markdown 内容偏短 {path.name}")
    if not re.search(r"(^|\n)#{1,6}\s+", text):
        warnings.append(f"{name}: Markdown 未发现标题 {path.name}")
    replacement_ratio = text.count("\ufffd") / max(1, len(text))
    if replacement_ratio > 0.01:
        issues.append(f"{name}: Markdown 可能存在编码乱码 {path.name}")
    if structured_images:
        from markdown_bundle import _links
        from urllib.parse import unquote, urlsplit
        refs = [unquote(urlsplit(target).path) for target, image in _links(text) if image]
    else:
        refs = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", text)
    for ref in refs:
        if ref.startswith("data:") or re.match(r"^[a-z]+://", ref):
            continue
        img_path = (path.parent / ref).resolve()
        if not img_path.exists():
            issues.append(f"{name}: Markdown 图片引用缺失 {ref}")


def check_json(name: str, path: Path, issues: list[str]) -> None:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        issues.append(f"{name}: JSON 无法解析 {path.name}: {exc}")
        return
    if not data:
        issues.append(f"{name}: JSON 是空壳 {path.name}")
    if isinstance(data, list) and len(data) == 0:
        issues.append(f"{name}: JSON 内容块数量为 0 {path.name}")


def check_csv(name: str, path: Path, issues: list[str], warnings: list[str]) -> None:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        issues.append(f"{name}: CSV 为空 {path.name}")
    elif len(rows) == 1:
        warnings.append(f"{name}: CSV 只有一行 {path.name}")


def render_quality_report(
    batch: Path,
    rows: list[dict[str, str]],
    issues: list[str],
    warnings: list[str],
) -> str:
    selected = [row for row in rows if row.get("primary_input")]
    review = [
        row
        for row in rows
        if row.get("quality_status") == "needs_manual_review"
        or row.get("processing_status") == "needs_manual_review"
        or row.get("processing_status") == "failed"
    ]
    lines = [
        "# 质量检查报告",
        "",
        f"- 批次：`{batch.name}`",
        f"- 检查时间：{dt.datetime.now().isoformat(timespec='seconds')}",
        f"- 本次记录：{len(rows)}",
        f"- 入库主输入：{len(selected)}",
        f"- 需人工复核：{len(review)}",
        f"- 严重问题：{len(issues)}",
        f"- 提醒：{len(warnings)}",
        "",
        "## 严重问题",
        "",
    ]
    lines.extend([f"- {issue}" for issue in issues] or ["- 无"])
    lines.extend(["", "## 提醒", ""])
    lines.extend([f"- {warning}" for warning in warnings] or ["- 无"])
    lines.extend(["", "## 需人工复核", ""])
    if review:
        for row in review:
            lines.append(f"- {row.get('file_name', '')}: {row.get('notes', '')}")
    else:
        lines.append("- 无")
    lines.append("")
    return "\n".join(lines)


def cmd_publish(args: argparse.Namespace) -> int:
    cfg = load_config(args.config)
    batch = args.batch.resolve()
    source = batch_dirs(batch)["kb"]
    if not source.exists():
        print("缺少 04_知识库输入文件，无法发布。", file=sys.stderr)
        return 2
    topic = safe_name(args.topic or infer_topic_from_batch(batch.name))
    date_part = dt.datetime.now().strftime("%Y%m%d")
    dest_name = f"{cfg.result_prefix}_{date_part}_{topic}_知识库输入文件"
    dest = unique_path(cfg.icloud_inbox / dest_name)
    shutil.copytree(source, dest)
    remove_junk_files(dest)
    print(f"已复制成果到: {dest}")
    return 0


def infer_topic_from_batch(batch_name: str) -> str:
    match = re.match(r"^\d{8}_\d{6}_(.+)$", batch_name)
    return match.group(1) if match else batch_name


def cmd_run(args: argparse.Namespace) -> int:
    cfg = load_config(args.config)
    topic = args.topic or cfg.default_topic
    candidates = select_candidates(cfg, args)
    if not candidates:
        print("没有找到需要处理的候选文件。")
        return 0
    batch_id = now_batch_id(topic)
    batch = cfg.local_transfer_station / batch_id
    prepare_args = argparse.Namespace(
        config=args.config,
        topic=topic,
        include=args.include,
        exclude=args.exclude,
        limit=args.limit,
    )
    original_now_batch_id = now_batch_id
    try:
        globals()["now_batch_id"] = lambda _topic: batch_id
        cmd_prepare(prepare_args)
    finally:
        globals()["now_batch_id"] = original_now_batch_id
    process_args = argparse.Namespace(config=args.config, batch=batch, force=False)
    cmd_process(process_args)
    check_args = argparse.Namespace(config=args.config, batch=batch)
    check_code = cmd_check(check_args)
    if check_code and not args.publish_with_issues:
        print("质量检查存在严重问题，已停止发布。可修复后单独运行 publish。")
        return check_code
    if args.skip_publish:
        print(f"已按要求跳过发布，批次目录: {batch}")
        return check_code
    publish_args = argparse.Namespace(config=args.config, batch=batch, topic=topic)
    return cmd_publish(publish_args)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="知识库入库半自动流程")
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG, help="配置文件路径")
    sub = parser.add_subparsers(dest="command", required=True)

    plan = sub.add_parser("plan", help="扫描 iCloud 投递箱，列出候选文件")
    plan.add_argument("--topic", default=None, help="本次主题，仅用于显示")
    plan.add_argument("--include", action="append", help="只选择路径中包含该片段的文件，可重复")
    plan.add_argument("--exclude", action="append", help="排除路径中包含该片段的文件，可重复")
    plan.add_argument("--limit", type=int, default=None, help="最多选择多少个候选文件")
    plan.add_argument("--verbose", action="store_true", help="显示候选文件明细")
    plan.set_defaults(func=cmd_plan)

    prepare = sub.add_parser("prepare", help="复制候选文件到本地中转站并建立批次")
    prepare.add_argument("--topic", default=None, help="本次批次主题")
    prepare.add_argument("--include", action="append", help="只选择路径中包含该片段的文件，可重复")
    prepare.add_argument("--exclude", action="append", help="排除路径中包含该片段的文件，可重复")
    prepare.add_argument("--limit", type=int, default=None, help="最多复制多少个候选文件")
    prepare.set_defaults(func=cmd_prepare)

    process = sub.add_parser("process", help="处理一个已准备好的批次")
    process.add_argument("--batch", type=Path, required=True, help="批次目录")
    process.add_argument("--force", action="store_true", help="重新处理已处理记录")
    process.set_defaults(func=cmd_process)

    check = sub.add_parser("check", help="检查知识库输入文件质量")
    check.add_argument("--batch", type=Path, required=True, help="批次目录")
    check.set_defaults(func=cmd_check)

    publish = sub.add_parser("publish", help="把成果复制回 iCloud 投递箱")
    publish.add_argument("--batch", type=Path, required=True, help="批次目录")
    publish.add_argument("--topic", default=None, help="成果目录主题名")
    publish.set_defaults(func=cmd_publish)

    run = sub.add_parser("run", help="串联执行 prepare/process/check/publish")
    run.add_argument("--topic", default=None, help="本次批次主题")
    run.add_argument("--include", action="append", help="只选择路径中包含该片段的文件，可重复")
    run.add_argument("--exclude", action="append", help="排除路径中包含该片段的文件，可重复")
    run.add_argument("--limit", type=int, default=None, help="最多处理多少个候选文件")
    run.add_argument(
        "--publish-with-issues",
        action="store_true",
        help="即使质量检查存在严重问题也复制成果回 iCloud",
    )
    run.add_argument(
        "--skip-publish",
        action="store_true",
        help="完成 prepare/process/check 后停止，不复制成果回 iCloud",
    )
    run.set_defaults(func=cmd_run)

    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
