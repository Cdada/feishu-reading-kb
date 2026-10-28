"""Bounded, offline TXT/OOXML extraction; originals are never rewritten.

Library readers handle the formats. The XML preflight rejects unsupported
active/ambiguous content and bounds expansion before those readers allocate.
All locations below describe document structure, never invented printed pages.
"""
from dataclasses import dataclass, field
from datetime import date, datetime, time
import hashlib
from html import escape
from io import BytesIO
from pathlib import PurePosixPath
import re
import zipfile

from defusedxml.ElementTree import fromstring
from defusedxml.common import DefusedXmlException
from xml.etree.ElementTree import ParseError
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries

from .contracts import ContractError

FORMATS = {'.txt', '.docx', '.xlsx'}
MAX_INPUT = 32 * 1024**2
MAX_EXPANDED = 64 * 1024**2
MAX_PART = 16 * 1024**2
MAX_TEXT = 1024 * 1024
MAX_CELLS = 50000
X = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
R = '{http://schemas.openxmlformats.org/package/2006/relationships}'


class DocumentProblem(ContractError):
    """Safe, actionable diagnostics without source excerpts or machine paths."""


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def plain(value):
    # Plain-text input must not become Markdown links, raw HTML or images.
    return re.sub(r'([\\`*_{}\[\]()#+.!|~-])', r'\\\1', escape(str(value), quote=False))


@dataclass
class ExtractedDocument:
    parts: list = field(default_factory=list)
    assets: dict = field(default_factory=dict)
    locations: list = field(default_factory=list)
    details: dict = field(default_factory=dict)
    length: int = 0

    def add(self, text, **location):
        text = text + '\n\n'
        if self.length + len(text) > MAX_TEXT or len(self.locations) >= 50000:
            raise DocumentProblem('转换正文超过完整保留上限，请拆分文档后再提交；没有截断正文。')
        if location:
            self.locations.append(dict(char_start=self.length, char_end=self.length+len(text),
                                       sha256=digest(text.encode()), **location))
        self.parts.append(text)
        self.length += len(text)

    @property
    def text(self):
        return ''.join(self.parts)

    def mapping(self, source_sha):
        return dict(reference='office-extraction/v1', source_sha256=source_sha,
                    output_sha256=digest(self.text.encode()), locations=self.locations,
                    details=self.details, physical_pages_available=False)


def package(raw, extension):
    """Read a bounded package in memory; no archive path is extracted to disk."""
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            infos = archive.infolist()
            names = [i.filename for i in infos]
            if (len(infos) > 2000 or len(names) != len(set(names))
                    or sum(i.file_size for i in infos) > MAX_EXPANDED):
                raise DocumentProblem('Office 压缩包重复条目或展开体积超限，请拆分或重新另存。')
            for i in infos:
                p = PurePosixPath(i.filename)
                if (p.is_absolute() or '..' in p.parts or '\\' in i.filename or ':' in i.filename
                        or i.flag_bits & 1 or i.file_size > MAX_PART
                        or i.file_size > max(1, i.compress_size)*500
                        or i.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED)):
                    raise DocumentProblem('Office 压缩包含不安全条目、加密内容或异常压缩比例。')
            blobs = {i.filename: archive.read(i) for i in infos if not i.is_dir()}
    except (zipfile.BadZipFile, RuntimeError, NotImplementedError) as exc:
        raise DocumentProblem('文件不是可读取的 DOCX/XLSX，请解除密码保护并重新另存。') from exc
    required = 'word/document.xml' if extension == '.docx' else 'xl/workbook.xml'
    if required not in blobs or '[Content_Types].xml' not in blobs:
        raise DocumentProblem('文件内容与扩展名不符，请另存为 DOCX 或 XLSX。')
    if (any('vbaproject' in n.lower() or '/embeddings/' in n.lower() for n in blobs)
            or b'macroEnabled' in blobs['[Content_Types].xml']):
        raise DocumentProblem('文件含宏或嵌入对象，当前请提交不含这些对象的副本。')
    xml = {}
    try:
        for name, value in blobs.items():
            if name.endswith(('.xml', '.rels')):
                xml[name] = fromstring(value, forbid_dtd=True)
                if name.endswith('.rels'):
                    for rel in xml[name]:
                        if rel.get('TargetMode') == 'External' and not rel.get('Type', '').endswith('/hyperlink'):
                            raise DocumentProblem('文件依赖外部图片或数据，请先将内容嵌入文件；不会联网补取。')
    except (DefusedXmlException, ParseError) as exc:
        raise DocumentProblem('Office XML 不安全或已损坏，请重新另存文件。') from exc
    return blobs, xml


def text_document(raw):
    encoding = 'utf-8-sig'
    if raw.startswith((b'\xff\xfe\x00\x00', b'\x00\x00\xfe\xff')):
        encoding = 'utf-32'
    elif raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        encoding = 'utf-16'
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError:
        if encoding != 'utf-8-sig':
            raise DocumentProblem('TXT 编码损坏，请另存为 UTF-8 后重新提交。') from None
        encoding = 'gb18030'
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError as exc:
            raise DocumentProblem('TXT 编码无法可靠读取，请另存为 UTF-8 后重新提交。') from exc
    if not text.strip() or '\ufffd' in text or any(ord(c) < 32 and c not in '\r\n\t' for c in text):
        raise DocumentProblem('TXT 为空、含替换乱码或二进制控制字符，请检查原文件。')
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    result = ExtractedDocument(details=dict(format='txt', encoding=encoding, newline='LF'))
    result.add('# TXT 正文\n\n> 系统说明：按 '+encoding+' 解码；正文按纯文本保留。')
    for number, line in enumerate(text.split('\n'), 1):
        result.add(plain(line), kind='text_line', line=number)
    return result


def retain_media(result, blobs, prefix):
    names = {}
    for name, value in blobs.items():
        if name.startswith(prefix+'/media/'):
            ext = PurePosixPath(name).suffix.lower()
            if ext not in {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tiff', '.webp'}:
                raise DocumentProblem('含暂不支持的矢量图或图片格式，请转为 PNG/JPEG 或 PDF 后提交。')
            path = 'images/'+digest(value)+ext
            names[name] = path
            result.assets[path] = value
    return names


def word_document(raw, blobs, xml):
    main = xml['word/document.xml']
    # These change visible text or introduce alternatives the reader cannot
    # faithfully linearize. Fail before presenting partial content as complete.
    unsupported = {'ins', 'del', 'moveFrom', 'moveTo', 'sdt', 'altChunk', 'txbxContent',
                   'object', 'fldSimple', 'fldChar', 'instrText'}
    if any(e.tag in {W+n for n in unsupported} or e.tag.endswith(('}oMath', '}AlternateContent', '}chart', '}relIds', '}pict')) for e in main.iter()):
        raise DocumentProblem('Word 含修订、域、文本框、公式或内容控件，请先接受修订/转为普通正文，或提交 PDF。')
    if not any((e.text or '').strip() for e in main.iter(W+'t')) and not any(e.tag == qn('a:blip') for e in main.iter()):
        raise DocumentProblem('Word 正文为空或只有未支持的对象，请提交可读取的正文。')
    document = Document(BytesIO(raw))
    out = ExtractedDocument(details=dict(format='docx', image_reading='not_read', pagination='unavailable'))
    media = retain_media(out, blobs, 'word')
    used = set()
    out.add('# Word 正文\n\n> 系统说明：按段落和表格顺序提取；未重建排版页码。图片原件保留，尚未识读；批注、自动列表编号和样式未还原。')

    def paragraph(p, in_table=False):
        parts = []
        for node in p._p.iter():
            if node.tag == qn('w:t'):
                parts.append(escape(node.text or '', quote=False) if in_table else plain(node.text or ''))
            elif node.tag == qn('w:tab'):
                parts.append('\t')
            elif node.tag in (qn('w:br'), qn('w:cr')):
                parts.append('\n')
            elif node.tag in (qn('w:footnoteReference'), qn('w:endnoteReference')):
                label = '脚注' if node.tag == qn('w:footnoteReference') else '尾注'
                parts.append('〔'+label+node.get(qn('w:id'), '')+'〕')
            elif node.tag == qn('a:blip'):
                rid = node.get(qn('r:embed'))
                if not rid or rid not in p.part.rels or p.part.rels[rid].is_external:
                    raise DocumentProblem('Word 图片关联不完整，请补全图片后重新提交。')
                name = str(p.part.rels[rid].target_part.partname).lstrip('/')
                if name not in media:
                    raise DocumentProblem('Word 图片未包含在受支持的内嵌图片中。')
                target = media[name]
                if in_table:
                    parts.append('【表内原图见附录，未识读】')
                else:
                    used.add(target)
                    parts.append('![原图，未识读]('+target+')')
        return ''.join(parts)

    def table(t, depth=0):
        if depth > 4 or len(t.rows)*len(t.columns) > MAX_CELLS:
            raise DocumentProblem('Word 表格过大或嵌套过深，请拆分后提交。')
        positions, objects = {}, {}
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells, row.grid_cols_before):
                key = cell._tc
                positions.setdefault(key, []).append((ri, ci)); objects[key] = cell
        origins, covered = {}, set()
        for key, coords in positions.items():
            rs = [p[0] for p in coords]; cs = [p[1] for p in coords]
            a, b, c, d = min(rs), max(rs), min(cs), max(cs)
            if len(coords) != (b-a+1)*(d-c+1):
                raise DocumentProblem('Word 合并表格结构不规则，请转为规则表格或 PDF。')
            content = []
            for child in objects[key].iter_inner_content():
                content.append(paragraph(child, in_table=True) if isinstance(child, Paragraph) else table(child, depth+1))
            origins[a, c] = (b-a+1, d-c+1, '<br>'.join(content).replace('\n', '<br>'))
            covered.update(coords)
        rows = []
        for ri, row in enumerate(t.rows):
            # Only an explicit Word repeat-header flag grants header semantics.
            flags = row._tr.xpath('./w:trPr/w:tblHeader')
            header = any(f.get(qn('w:val'), 'true') not in ('0','false','off') for f in flags)
            tag = 'th' if header else 'td'
            cells = []
            for ci in range(len(t.columns)):
                if (ri, ci) in origins:
                    h, w, value = origins[ri, ci]
                    cells.append(f'<{tag} rowspan="{h}" colspan="{w}">{value}</{tag}>')
                elif (ri, ci) not in covered:
                    cells.append('<td></td>')
            rows.append('<tr>'+''.join(cells)+'</tr>')
        return '<table>\n'+'\n'.join(rows)+'\n</table>'

    for index, block in enumerate(document.iter_inner_content(), 1):
        if isinstance(block, Paragraph):
            out.add(paragraph(block), kind='word_paragraph', block=index)
        else:
            out.add(table(block), kind='word_table', block=index)
    for name in ('word/footnotes.xml', 'word/endnotes.xml'):
        if name in xml:
            label = '脚注' if 'footnotes' in name else '尾注'
            for note in xml[name]:
                if note.get(W+'type') in ('separator', 'continuationSeparator'):
                    continue
                # Footnote bodies with drawings/fields require a richer reader.
                if any(e.tag in (W+'drawing', W+'tbl', W+'fldChar') for e in note.iter()):
                    raise DocumentProblem('Word 注释含复杂对象，请提交 PDF 以保留完整注释。')
                value = ''.join((e.text or '') if e.tag == W+'t' else '\n' if e.tag in (W+'p', W+'br', W+'cr') else '\t' if e.tag == W+'tab' else '' for e in note.iter())
                out.add('## '+label+note.get(W+'id', '')+'\n\n'+plain(value), kind='word_note', part=name, note_id=note.get(W+'id'))
    # Retain header/footer text once, explicitly separate from substantive body.
    for name, element in xml.items():
        if re.fullmatch(r'word/(header|footer)\d+\.xml', name):
            value = '\n'.join(''.join(t.text or '' for t in p.iter(W+'t')) for p in element.iter(W+'p'))
            if value.strip():
                out.add('## 页眉/页脚附录（版面信息）\n\n'+plain(value), kind='word_furniture', part=name)
    for target in sorted(set(media.values())-used):
        out.add('![补充原图，未识读，位置未还原]('+target+')', kind='unlocated_image', target=target)
    return out


def spreadsheet(raw, blobs, xml):
    sheets = {}
    total = 0
    for name, root in xml.items():
        if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', name):
            cells = list(root.iter(X+'c'))
            rows, cols = 0, 0
            seen = set()
            for cell in cells:
                if cell.get('r') in seen:
                    raise DocumentProblem('Excel 含重复单元格坐标，请修复原表。')
                seen.add(cell.get('r'))
                ri, ci = coordinate_to_tuple(cell.get('r', ''))
                if ri < 1 or ci < 1:
                    raise DocumentProblem('Excel 单元格坐标无效。')
                rows, cols = max(rows, ri), max(cols, ci)
            merges = [e.get('ref') for e in root.iter(X+'mergeCell')]
            for ref in merges:
                a, b, ci, ri = range_boundaries(ref)
                if not all(isinstance(n, int) and n > 0 for n in (a,b,ci,ri)) or a > ci or b > ri:
                    raise DocumentProblem('Excel 合并范围无效。')
                rows, cols = max(rows, ri), max(cols, ci)
            total += rows*cols
            if rows > 10000 or cols > 256 or total > MAX_CELLS or len(sheets) >= 100:
                raise DocumentProblem('Excel 数据区域超过 5 万格、256 列或 100 个工作表，请拆分；不会只取前几行。')
            sheets[name] = (rows, cols, merges, root)
    if not total:
        raise DocumentProblem('Excel 工作簿没有可读取的单元格内容。')
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    out = ExtractedDocument(details=dict(format='xlsx', formulas='not_calculated', cached_values='not_used',
                                         hidden_cells='included_and_marked', formatting='stored_value_and_number_format',
                                         charts='not_read', sheet_count=len(workbook.worksheets)))
    media = retain_media(out, blobs, 'xl')
    out.add('# Excel 工作簿\n\n> 系统说明：完整提取有数据的工作表范围，A/B 与行号是系统坐标，不代表原表业务表头。'
            '公式仅保留表达式，不计算、不使用缓存结果；数值附带原始数字格式，不能当作已渲染的显示值。'
            '隐藏行列一并保留并标注。图表、图像、批注及排版效果未识读，请查看原件。')
    try:
        for sheet_number, sheet in enumerate(workbook.worksheets, 1):
            path = sheet._worksheet_path.lstrip('/')
            if path not in sheets:
                raise DocumentProblem('Excel 工作表关系不符合支持范围，请重新另存。')
            height, width, merges, root = sheets[path]
            hidden_rows = [e.get('r') for e in root.iter(X+'row') if e.get('hidden') in ('1', 'true')]
            hidden_cols = [e.get('min')+':'+e.get('max') for e in root.iter(X+'col') if e.get('hidden') in ('1', 'true')]
            out.add('## 工作表：'+plain(sheet.title)+'\n\n状态：'+sheet.sheet_state+
                    '；隐藏行：'+(','.join(hidden_rows) or '无')+'；隐藏列范围：'+(','.join(hidden_cols) or '无'),
                    kind='worksheet_heading', sheet=sheet.title)
            if not height or not width:
                out.add('（空工作表）', kind='worksheet_empty', sheet=sheet.title)
                continue
            merged, occupied = {}, set()
            for ref in merges:
                a, b, c, d = range_boundaries(ref)
                coords = {(ri, ci) for ri in range(b, d+1) for ci in range(a, c+1)}
                if occupied & coords:
                    raise DocumentProblem('Excel 合并单元格相互重叠，请修复原表。')
                occupied.update(coords); merged[b, a] = (d-b+1, c-a+1)
            table = ['<table>', '<tr><th>行号（系统坐标）</th>'+''.join('<th>'+get_column_letter(c)+'（系统列号）</th>' for c in range(1,width+1))+'</tr>']
            for ri, row in enumerate(sheet.iter_rows(min_row=1, max_row=height, min_col=1, max_col=width), 1):
                values = ['<td>'+str(ri)+'</td>']
                for ci, cell in enumerate(row, 1):
                    if (ri, ci) in occupied and (ri, ci) not in merged:
                        if cell.value is not None:
                            raise DocumentProblem('Excel 合并覆盖区域仍含值，请检查原表；没有丢弃这些值。')
                        continue
                    v = cell.value
                    if cell.data_type == 'f':
                        if not isinstance(v, str):
                            raise DocumentProblem('Excel 含数组或数据表公式，请先另存为值副本。')
                        text = '【公式未计算：'+v+'】'
                    elif isinstance(v, (datetime, date, time)):
                        text = v.isoformat()+'（数字格式：'+cell.number_format+'）'
                    elif isinstance(v, (int, float)) and not isinstance(v, bool) and cell.number_format != 'General':
                        text = str(v)+'（数字格式：'+cell.number_format+'）'
                    else:
                        text = '' if v is None else str(v)
                    h, w = merged.get((ri, ci), (1, 1))
                    values.append(f'<td rowspan="{h}" colspan="{w}">'+escape(text).replace('\n','<br>')+'</td>')
                table.append('<tr>'+''.join(values)+'</tr>')
            table.append('</table>')
            caption = '表 '+str(sheet_number)+' 工作表：'+plain(sheet.title)+'（系统生成定位）\n\n说明：A/B 与行号是坐标，非业务表头；公式未计算；数字格式未渲染。\n\n'
            out.add(caption+'\n'.join(table), kind='worksheet_cells', sheet=sheet.title,
                    cell_range='A1:'+get_column_letter(width)+str(height), merged_ranges=merges)
        for target in sorted(set(media.values())):
            out.add('![工作簿原图，未识读，位置未还原]('+target+')', kind='unlocated_image', target=target)
    finally:
        workbook.close()
    return out


def extract_document(raw, extension):
    if not isinstance(raw, bytes) or len(raw) > MAX_INPUT:
        raise DocumentProblem('文件超过 32 MiB，请拆分后再提交。')
    if extension not in FORMATS:
        raise DocumentProblem('当前支持 TXT、DOCX、XLSX；旧版 DOC/XLS 请先另存为新版格式。')
    if extension == '.txt':
        return text_document(raw)
    blobs, xml = package(raw, extension)
    try:
        result = word_document(raw, blobs, xml) if extension == '.docx' else spreadsheet(raw, blobs, xml)
    except DocumentProblem:
        raise
    except (KeyError, ValueError, IndexError, TypeError, OSError) as exc:
        raise DocumentProblem('Office 内容结构损坏或超出当前解析范围，请重新另存或转为 PDF。') from exc
    if not result.locations:
        raise DocumentProblem('文件中没有可归档的正文，请检查原件。')
    return result
