"""Synthetic documents only; no business model, Office app or HTTP calls."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook

from kb_custody_v21.office_documents import extract_document, DocumentProblem, package
from kb_custody_v21.document_formats import daily_formats, OFFICE_PROFILE, OFFICE_FORMATS
from kb_custody_v21.table_evidence import inventory
from kb_custody_v21.f01 import encoded
import test_conversion_v21 as conversion_fixture
import test_pilot_three_formats as guide_fixture


def word(*, header=True, image=False):
    doc = Document(); doc.add_paragraph('合成前文：资料必须按原件顺序保留。')
    table = doc.add_table(rows=3, cols=3)
    for row, values in zip(table.rows, [('地区','面积','比例'),('合成甲区','12.5','10%'),('合成乙区','0','—')]):
        for cell, value in zip(row.cells, values): cell.text = value
    if header:
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1,0).merge(table.cell(2,0))
    if image:
        from test_markdown_bundle import png
        doc.add_picture(BytesIO(png()))
    doc.add_paragraph('合成后文：不能凭空补充数据。')
    stream=BytesIO(); doc.save(stream); return stream.getvalue()


def excel(*, rows=3):
    book=Workbook(); sheet=book.active; sheet.title='合成资料'
    sheet.append(['地区','面积','比例'])
    sheet.append(['合成甲区',12.5,0.1]); sheet['C2'].number_format='0%'
    sheet.append(['合成乙区',0,'=SUM(B2:B3)'])
    for i in range(4, rows+1): sheet.append(['合成行'+str(i),i,None])
    sheet.row_dimensions[3].hidden=True
    extra=book.create_sheet('隐藏记录'); extra.sheet_state='hidden'; extra['A1']='仍须完整保留'
    stream=BytesIO(); book.save(stream); return stream.getvalue()


def amend(raw, name, mutate):
    stream=BytesIO()
    with ZipFile(BytesIO(raw)) as src, ZipFile(stream,'w',ZIP_DEFLATED) as dest:
        for item in src.infolist():
            data=src.read(item.filename)
            dest.writestr(item.filename, mutate(data) if item.filename==name else data)
    return stream.getvalue()


class OfficeExtractionTests(unittest.TestCase):
    def test_package_root_cannot_redirect_around_the_validated_body(self):
        for raw, ext, old, new in ((word(), '.docx', b'word/document.xml', b'word/alternate.xml'),
                                   (excel(), '.xlsx', b'xl/workbook.xml', b'xl/alternate.xml')):
            changed=amend(raw,'_rels/.rels',lambda value:value.replace(old,new))
            with self.subTest(ext=ext),self.assertRaisesRegex(DocumentProblem,'主文档关系'):
                extract_document(changed,ext)

    def test_word_unsupported_visible_content_is_never_silently_dropped(self):
        for tag in ('customXml', 'sym', 'vanish'):
            raw=amend(word(),'word/document.xml',lambda b:b.replace(b'<w:body>',b'<w:body><w:'+tag.encode()+b'/>',1))
            with self.subTest(tag=tag),self.assertRaisesRegex(DocumentProblem,'Word 含'):
                extract_document(raw,'.docx')

    def test_excel_row_and_cell_coordinates_must_agree(self):
        raw=amend(excel(),'xl/worksheets/sheet1.xml',lambda b:b.replace(b'r="B2"',b'r="B1"',1))
        with self.assertRaisesRegex(DocumentProblem,'坐标'):
            extract_document(raw,'.xlsx')

    def test_word_missing_grid_never_becomes_an_empty_table(self):
        import re
        raw=amend(word(),'word/document.xml',lambda b:re.sub(rb'<w:tblGrid>.*?</w:tblGrid>',b'<w:tblGrid/>',b))
        with self.assertRaisesRegex(DocumentProblem,'列结构'):
            extract_document(raw,'.docx')

    def test_txt_encodings_and_literal_markup(self):
        for encoding in ('utf-8','utf-8-sig','utf-16','utf-32','gb18030'):
            with self.subTest(encoding=encoding):
                result=extract_document('简体中文与繁體字\r\n<table>不是表格</table>\n![x](remote.png)'.encode(encoding),'.txt')
                self.assertIn('简体中文与繁體字',result.text)
                self.assertNotIn('\r',result.text)
                self.assertEqual(inventory(result.text)['tables'],[])
                self.assertEqual(result.assets,{})
                self.assertEqual(len(result.locations),3)

    def test_txt_rejects_binary_empty_broken_and_replacement(self):
        for raw in (b'\x00abc',b'   ',b'\xff',b'\xff\xfeA','\ufffd错误'.encode()):
            with self.subTest(raw=raw), self.assertRaises(DocumentProblem): extract_document(raw,'.txt')

    def test_word_order_merge_explicit_header_and_no_invented_pages(self):
        result=extract_document(word(),'.docx')
        self.assertLess(result.text.index('合成前文'),result.text.index('<table>'))
        self.assertLess(result.text.index('</table>'),result.text.index('合成后文'))
        table=inventory(result.text)['tables'][0]
        self.assertTrue(table['qa_usable']); self.assertEqual(table['issues'],[])
        self.assertIn('12.5', [c['text'] for c in table['cells']])
        self.assertTrue(any(c['rowspan']==2 for c in table['cells']))
        self.assertEqual(table['physical_pages'],[])
        self.assertFalse(result.mapping('a'*64)['physical_pages_available'])

    def test_word_without_explicit_header_does_not_invent_one(self):
        table=inventory(extract_document(word(header=False),'.docx').text)['tables'][0]
        self.assertEqual(table['header_status'],'not_marked'); self.assertFalse(table['qa_usable'])

    def test_word_images_kept_and_not_read(self):
        from test_markdown_bundle import png
        result=extract_document(word(image=True),'.docx')
        self.assertIn(png(),result.assets.values()); self.assertIn('原图，未识读',result.text)

    def test_word_revision_and_textbox_fail_before_partial_result(self):
        for tag in ('ins','txbxContent','sdt','fldChar'):
            raw=amend(word(),'word/document.xml',lambda b:b.replace(b'<w:body>',b'<w:body><w:'+tag.encode()+b'/>',1))
            with self.subTest(tag=tag),self.assertRaisesRegex(DocumentProblem,'Word 含'): extract_document(raw,'.docx')

    def test_excel_complete_hidden_and_formula_boundaries(self):
        result=extract_document(excel(rows=125),'.xlsx')
        self.assertIn('合成行125',result.text);self.assertIn('仍须完整保留',result.text)
        self.assertIn('隐藏行：3',result.text);self.assertIn('状态：hidden',result.text)
        self.assertIn('【公式未计算：=SUM(B2:B3)】',result.text)
        self.assertIn('0.1（数字格式：0%）',result.text)
        self.assertEqual([p['sheet'] for p in result.locations if p['kind']=='worksheet_cells'],['合成资料','隐藏记录'])
        tables=inventory(result.text)['tables'];self.assertEqual(len(tables),2)
        self.assertEqual(tables[0]['rows'],126);self.assertEqual(tables[0]['issues'],[])

    def test_excel_keeps_merged_cells_empty_and_zero_distinct(self):
        book=Workbook();s=book.active;s['A1']='合成分类';s.merge_cells('A1:B2');s['C1']=0;s['C2']='—'
        b=BytesIO();book.save(b)
        result=extract_document(b.getvalue(),'.xlsx');t=inventory(result.text)['tables'][0]
        self.assertEqual(t['issues'],[])
        self.assertTrue(any(c['rowspan']==2 and c['colspan']==2 for c in t['cells']))
        self.assertIn('0',[c['text'] for c in t['cells']])

    def test_excel_huge_sparse_rectangle_rejected_before_library(self):
        book=Workbook();book.active['XFD1048576']='合成极远单元格';b=BytesIO();book.save(b)
        with patch('kb_custody_v21.office_documents.load_workbook',side_effect=AssertionError('must reject first')):
            with self.assertRaisesRegex(DocumentProblem,'超过'):extract_document(b.getvalue(),'.xlsx')

    def test_wrong_extension_password_container_and_legacy(self):
        for raw,ext in [(word(),'.xlsx'),(b'encrypted','.docx'),(b'old','.doc'),(b'old','.xls')]:
            with self.subTest(ext=ext),self.assertRaises(DocumentProblem):extract_document(raw,ext)

    def test_package_rejects_paths_duplicates_macros_dtd_and_external_media(self):
        bads=[]
        for name in ('../bad.xml','xl/vbaProject.bin','word/embeddings/object.bin'):
            b=BytesIO(word())
            with ZipFile(b,'a') as z:z.writestr(name,b'synthetic')
            bads.append(b.getvalue())
        bads.append(amend(word(),'word/document.xml',lambda b:b.replace(b'?>',b'?> <!DOCTYPE x [<!ENTITY x "expansion">]>',1)))
        bads.append(amend(word(),'word/_rels/document.xml.rels',lambda b:b.replace(b'</Relationships>',b'<Relationship Id="rId99" Type="image" Target="https://example.invalid/a.png" TargetMode="External"/></Relationships>')))
        for raw in bads:
            with self.assertRaises(DocumentProblem):extract_document(raw,'.docx')

    def test_structure_offsets_are_bound_to_exact_text(self):
        import hashlib
        for raw,ext in [(word(),'.docx'),(excel(),'.xlsx')]:
            r=extract_document(raw,ext)
            for loc in r.locations:
                self.assertEqual(loc['sha256'],hashlib.sha256(r.text[loc['char_start']:loc['char_end']].encode()).hexdigest())

    def test_versioned_profile_only_enables_new_formats(self):
        with tempfile.TemporaryDirectory() as root:
            p=Path(root)/'profile.json';p.write_text(json.dumps(dict(reference='P33-pdf-guide-v1',formats=['.md','.pdf'])))
            self.assertNotIn('.txt',daily_formats(dict(guide_plan=str(p),pdf={})))
            p.write_text(json.dumps(dict(reference=OFFICE_PROFILE,formats=OFFICE_FORMATS)))
            self.assertEqual(daily_formats(dict(guide_plan=str(p),pdf={})),OFFICE_FORMATS)
            self.assertNotIn('.pdf',daily_formats(dict(guide_plan=str(p))))


class OfficePipelineTests(unittest.TestCase):
    def test_real_worker_retains_original_and_structured_map(self):
        for raw,ext in [('合成文字资料。正文保持完整，不能虚构其来源与页码。'.encode('gb18030'),'.txt'),(word(),'.docx'),(excel(),'.xlsx')]:
            with self.subTest(ext=ext):
                f=conversion_fixture.ConversionTests();f.setUp();self.addCleanup(f.doCleanups)
                eid=f.receive(raw,received_filename='synthetic'+ext)
                _,mapping=f.convert(eid)
                self.assertEqual(mapping['locator_precision'],'document_structure')
                self.assertEqual(mapping['document_structure']['source_sha256'],f.registry.entry(eid)['sha256'])
                self.assertEqual(mapping['page_mappings'],[])
                original=f.runtime.root/f.registry.entry(eid)['relative_path']
                self.assertEqual(original.read_bytes(),raw)

    def test_new_formats_reach_preview_but_never_auto_confirm(self):
        for raw,ext in [('合成资料内容用于测试正文保留，不代表真实事实。'.encode(),'.txt'),(word(image=True),'.docx'),(excel(),'.xlsx')]:
            with self.subTest(ext=ext):
                f=guide_fixture.ThreeFormatGuideTests();f.setUp();self.addCleanup(f.doCleanups)
                f.g.plan.update(reference=OFFICE_PROFILE,formats=OFFICE_FORMATS,max_characters=20000,max_body_bytes=60000)
                f.g.f.approval['inbox']['formats']=OFFICE_FORMATS
                f.g.f.approval['inbox']['max_asset_bytes']=1024*1024
                f.g.save()
                eid=f.accept('synthetic'+ext,raw)
                result=f.g.guide(f.g.pipeline,eid)
                self.assertTrue(Path(result['preview_path']).is_file())
                self.assertEqual(f.g.wire.calls,1)
                self.assertFalse(f.g.registry.rows('SELECT * FROM publication_confirmations'))



class OfficeDailyFlowTests(unittest.TestCase):
    def test_owner_confirmation_to_reading_library_for_each_format(self):
        import test_pilot_inbox_guide
        import test_reading_flow
        from test_pilot_download import download_response
        old_accept=test_pilot_inbox_guide.InboxGuideTests.accept
        for raw,ext in [('合成材料正文用于验证完整整理，确认后才写入读书测试库。'.encode(),'.txt'),(word(),'.docx'),(excel(),'.xlsx')]:
            with self.subTest(ext=ext):
                def accept(g, *, convert=True):
                    g.plan.update(reference=OFFICE_PROFILE,formats=OFFICE_FORMATS,max_characters=20000,max_body_bytes=60000,
                        policy=dict(reference='P27-preserve-originals-v1',reading='not_requested',max_images=100,max_total_bytes=10_000_000))
                    g.f.approval['inbox'].update(formats=OFFICE_FORMATS,max_asset_bytes=1024*1024)
                    g.f.message['body']['content']=encoded(dict(file_key='file_synthetic',file_name='synthetic'+ext))
                    original=g.f.runner.side_effect
                    g.f.runner.side_effect=lambda c:download_response(c,raw) if c['kind']=='download' else original(c)
                    g.save()
                    return old_accept(g,convert=convert)
                with patch.object(test_pilot_inbox_guide.InboxGuideTests,'accept',accept):
                    f=test_reading_flow.ReadingFlowTests();f.setUp();self.addCleanup(f.doCleanups)
                document=f.f.preview()
                self.assertEqual(document['phase'],'awaiting_confirmation')
                self.assertFalse(f.a.reading.root.exists())
                f.f.owner_reply(document);f.b.advance_light(document);f.finish()
                self.assertEqual(f.document()['phase'],'notified')
                group=f.a.reading.root/f.document()['data']['reading_group']
                self.assertEqual((group/('入库依据/原件/source'+ext)).read_bytes(),raw)
                self.assertTrue((group/'综合读书笔记.md').is_file())
                self.assertEqual([r['purpose'] for r in f.wire.calls],['organize','review'])

    def test_mentions_use_same_versioned_format_allowlist(self):
        from test_daily_archive import ArchiveTests
        for ext in ('.txt','.docx','.xlsx'):
            f=ArchiveTests();f.setUp();self.addCleanup(f.doCleanups)
            profile=f.root/'guide.json';profile.write_text(encoded(dict(reference=OFFICE_PROFILE,formats=OFFICE_FORMATS)))
            f.a.config['guide_plan']=str(profile)
            body=json.loads(f.items['om_file']['body']['content']);body['file_name']='synthetic'+ext
            f.items['om_file']['body']['content']=encoded(body)
            f.persist();f.flow.drain()
            self.assertEqual(f.store.document('om_file')['phase'],'registered')

    def test_image_rewrite_preserves_all_structure_offsets(self):
        from kb_custody_v21.image_links import rewrite_images, image_spans
        from kb_custody_v21.office_documents import digest
        result=extract_document(word(image=True),'.docx');raw=result.text.encode()
        mapping=dict(extractor='custody-office-documents',output_sha256=digest(raw),transformations=[],document_structure=result.mapping('a'*64))
        replacements={s['original_path']:dict(asset_id='asset_synthetic',local_path='images/new-name.png',representation_asset_id='asset_local') for s in image_spans(result.text)}
        out,revised=rewrite_images(raw,mapping,replacements)
        self.assertNotEqual(out,raw)
        for loc in revised['document_structure']['locations']:
            self.assertEqual(loc['sha256'],digest(out.decode()[loc['char_start']:loc['char_end']].encode()))
        self.assertEqual(revised['document_structure']['output_sha256'],digest(out))

    def test_table_cell_image_is_retained_as_explicit_unlocated_appendix(self):
        from test_markdown_bundle import png
        from kb_custody_v21.image_links import image_spans
        doc=Document();doc.add_paragraph('合成表内图片')
        doc.add_table(rows=1,cols=1).cell(0,0).paragraphs[0].add_run().add_picture(BytesIO(png()))
        b=BytesIO();doc.save(b);result=extract_document(b.getvalue(),'.docx')
        self.assertEqual(len(image_spans(result.text)),1)
        self.assertIn('表内原图见附录',result.text)

    def test_excel_last_row_retrieval_keeps_sheet_and_coordinates(self):
        from kb_custody_v21.qa_index import QAIndex
        from test_qa_index import source, Provider
        result=extract_document(excel(rows=15),'.xlsx')
        with tempfile.TemporaryDirectory() as root:
            item=source('spreadsheet',result.text)
            index=QAIndex(Path(root).resolve()/'index.sqlite3')
            index.sync(Provider(item))
            hits=index.search('合成行15')
            self.assertTrue(hits)
            self.assertIn('合成行15',hits[0]['text'])
            self.assertIn('工作表：合成资料',hits[0]['text'])
            self.assertTrue(hits[0]['table_context'])


if __name__ == '__main__':
    unittest.main()
